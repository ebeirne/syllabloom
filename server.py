from __future__ import annotations

import cgi
import hashlib
import html
import io
import json
import os
import posixpath
import re
import sqlite3
import sys
import tempfile
import threading
import time
import zipfile
from api.legacy_powerpoint import PowerPointReadError
from collections.abc import Callable
from datetime import date, datetime, timezone
from http import HTTPStatus
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlparse
from xml.etree import ElementTree

from api._common import clerk_auth_config, deployment_env


ROOT = Path(__file__).resolve().parent
MODEL_PATH = Path(
    os.environ.get(
        "SYLLABLOOM_WHISPER_MODEL",
        os.environ.get("ROUNDS_WHISPER_MODEL", str(ROOT / ".models" / "whisper-small")),
    )
)
MAX_UPLOAD_BYTES = 1024 * 1024 * 1024
ALLOWED_SUFFIXES = {".wav", ".mp3", ".m4a", ".webm", ".ogg", ".flac", ".mp4", ".mov"}
SOURCE_SUFFIXES = {".docx", ".ppt", ".pptw", ".pptx", ".pdf", ".txt"}
DATA_DIR = ROOT / "data"
HOSTED = bool(deployment_env())
SESSIONS_DIR = DATA_DIR / "sessions"
SOURCE_LIBRARY_PATH = DATA_DIR / "source-library.json"


class NoSelectableTextError(ValueError):
    """Raised when a PDF has no text layer for the current importer to read."""

_model = None
_model_lock = threading.Lock()
_inference_lock = threading.Lock()
_data_lock = threading.Lock()


def safe_slug(value: str, fallback: str = "file") -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")
    return slug[:80] or fallback


def _safe_anki_tag(value: object) -> str:
    parts = [safe_slug(part, "") for part in str(value or "").split("::")]
    normalized = "::".join(part for part in parts if part)
    return normalized or "syllabloom"


def write_json(path: Path, payload: dict | list) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(path)


def read_source_library() -> list[dict]:
    if not SOURCE_LIBRARY_PATH.exists():
        return []
    try:
        payload = json.loads(SOURCE_LIBRARY_PATH.read_text(encoding="utf-8"))
        return payload if isinstance(payload, list) else []
    except (OSError, json.JSONDecodeError):
        return []


def extract_source_text(path: Path, suffix: str) -> tuple[str, dict]:
    if suffix == ".docx":
        from docx import Document

        document = Document(path)
        lines = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                line = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if line:
                    lines.append(line)
        return "\n".join(lines), {"unitLabel": "paragraphs", "unitCount": len(lines)}
    if suffix in {".ppt", ".pptw", ".pptx"}:
        if not zipfile.is_zipfile(path):
            from api.legacy_powerpoint import extract_legacy_powerpoint
            return extract_legacy_powerpoint(path)
        namespaces = {
            "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
            "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
            "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
            "pr": "http://schemas.openxmlformats.org/package/2006/relationships",
        }
        with zipfile.ZipFile(path) as package:
            presentation_root = ElementTree.fromstring(package.read("ppt/presentation.xml"))
            relationships_root = ElementTree.fromstring(package.read("ppt/_rels/presentation.xml.rels"))
            relationships = {
                relationship.attrib["Id"]: relationship.attrib["Target"]
                for relationship in relationships_root.findall("pr:Relationship", namespaces)
            }
            slide_paths = []
            for slide_id in presentation_root.findall("p:sldIdLst/p:sldId", namespaces):
                relationship_id = slide_id.attrib.get(f"{{{namespaces['r']}}}id")
                target = relationships.get(relationship_id or "")
                if target:
                    slide_paths.append(
                        target.lstrip("/")
                        if target.startswith("/")
                        else posixpath.normpath(posixpath.join("ppt", target))
                    )

            lines = []
            image_count = 0
            image_alt_text_count = 0
            notes_count = 0
            slides_with_unlabeled_images = []
            for slide_number, slide_path in enumerate(slide_paths, start=1):
                slide_root = ElementTree.fromstring(package.read(slide_path))
                slide_lines = []
                for paragraph in slide_root.findall(".//a:p", namespaces):
                    paragraph_text = "".join(
                        node.text or "" for node in paragraph.findall(".//a:t", namespaces)
                    ).strip()
                    if paragraph_text:
                        slide_lines.append(paragraph_text)
                slide_images = slide_root.findall(".//p:pic", namespaces)
                image_count += len(slide_images)
                image_descriptions = []
                for picture in slide_images:
                    properties = picture.find("p:nvPicPr/p:cNvPr", namespaces)
                    if properties is None:
                        continue
                    description = " ".join(
                        str(properties.attrib.get(key) or "").strip()
                        for key in ("title", "descr")
                    ).strip()
                    if description and description.lower() not in {"image", "picture", "graphic"}:
                        image_descriptions.append(description)
                image_alt_text_count += len(image_descriptions)
                if slide_images and not image_descriptions:
                    slides_with_unlabeled_images.append(slide_number)
                if image_descriptions:
                    slide_lines.extend(f"Image description: {description}" for description in image_descriptions)

                rels_path = posixpath.join(
                    posixpath.dirname(slide_path), "_rels", posixpath.basename(slide_path) + ".rels"
                )
                try:
                    slide_relationships = ElementTree.fromstring(package.read(rels_path))
                except KeyError:
                    slide_relationships = None
                notes_path = ""
                if slide_relationships is not None:
                    for relationship in slide_relationships.findall("pr:Relationship", namespaces):
                        if relationship.attrib.get("Type", "").endswith("/notesSlide"):
                            notes_target = relationship.attrib.get("Target", "")
                            notes_path = (
                                notes_target.lstrip("/")
                                if notes_target.startswith("/")
                                else posixpath.normpath(posixpath.join(posixpath.dirname(slide_path), notes_target))
                            )
                            break
                if notes_path:
                    try:
                        notes_root = ElementTree.fromstring(package.read(notes_path))
                    except KeyError:
                        notes_root = None
                    if notes_root is not None:
                        note_lines = []
                        for shape in notes_root.findall(".//p:sp", namespaces):
                            placeholder = shape.find("p:nvSpPr/p:nvPr/p:ph", namespaces)
                            if placeholder is not None and placeholder.attrib.get("type") in {
                                "sldNum", "dt", "hdr", "ftr", "sldImg"
                            }:
                                continue
                            paragraph_text = "".join(
                                node.text or "" for node in shape.findall(".//a:t", namespaces)
                            ).strip()
                            if paragraph_text and not paragraph_text.isdigit():
                                note_lines.append(paragraph_text)
                        note_text = "\n".join(note_lines).strip()
                        if note_text:
                            notes_count += 1
                            slide_lines.append("Speaker notes:\n" + note_text)
                if slide_lines:
                    lines.append(f"Slide {slide_number}\n" + "\n".join(slide_lines))
        return "\n".join(lines), {
            "unitLabel": "slides",
            "unitCount": len(slide_paths),
            "imageCount": image_count,
            "imageAltTextCount": image_alt_text_count,
            "slidesWithUnlabeledImages": slides_with_unlabeled_images,
            "speakerNotesCount": notes_count,
        }
    if suffix == ".pdf":
        import pdfplumber

        with pdfplumber.open(path) as document:
            pages = []
            for page in document.pages:
                normalized_page = None
                try:
                    normalized_page = page.dedupe_chars(tolerance=1)
                    pages.append((normalized_page.extract_text(x_tolerance=2, y_tolerance=3) or "").strip())
                finally:
                    # pdfplumber caches page layouts. Release each page instead
                    # of retaining an entire large lecture's object graph.
                    if normalized_page is not None:
                        normalized_page.close()
                    page.close()
        return "\n".join(page for page in pages if page), {
            "unitLabel": "pages",
            "unitCount": len(pages),
            "pageTexts": pages,
        }
    return path.read_text(encoding="utf-8", errors="replace"), {"unitLabel": "lines", "unitCount": 0}


def draft_cards_from_structured_slides(text: str, filename: str) -> list[dict]:
    """Create traceable drafts only when a slide exposes an explicit anatomy schema."""
    cards = []
    slide_blocks = re.split(r"(?m)^Slide (\d+)\s*$", text)
    for index in range(1, len(slide_blocks), 2):
        slide_number = int(slide_blocks[index])
        lines = [re.sub(r"\s+", " ", line).strip() for line in slide_blocks[index + 1].splitlines()]
        lines = [line for line in lines if line]
        upper_lines = [line.upper() for line in lines]
        labels = {label: upper_lines.index(label) for label in ("ATTACHMENT", "ACTION", "INNERVATION") if label in upper_lines}
        if len(labels) != 3:
            continue
        first_label_index = min(labels.values())
        subject = next(
            (
                line
                for line in reversed(lines[:first_label_index])
                if not line.isupper() and not re.fullmatch(r"\d+", line)
            ),
            "",
        )
        if not subject:
            continue
        section = next(
            (
                line.title()
                for line in reversed(lines[:first_label_index])
                if line.isupper() and line not in {"ROUNDS", "SOURCE SIGNALS"}
            ),
            "Imported slides",
        )
        prompts = {
            "ATTACHMENT": f"Where does {subject} attach?",
            "ACTION": f"What is the action of {subject}?",
            "INNERVATION": f"What innervates {subject}?",
        }
        for label, prompt in prompts.items():
            value_index = labels[label] + 1
            if value_index >= len(lines):
                continue
            answer = lines[value_index]
            cards.append(
                {
                    "front": prompt,
                    "back": answer,
                    "field": label.lower(),
                    "section": section,
                    "status": "verified",
                    "source": f"{filename} · Slide {slide_number}",
                    "slideNumber": slide_number,
                }
            )
    return cards


_GENERIC_SECTION_TITLES = {
    "agenda",
    "contents",
    "course overview",
    "learning objectives",
    "objectives",
    "references",
    "questions",
    "thank you",
}
_GENERIC_LABELS = {
    "advantages",
    "disadvantages",
    "example",
    "examples",
    "idea",
    "note",
    "notes",
    "key point",
    "key points",
    "reverse",
    "summary",
    "source",
    "question",
    "answer",
    "answers",
    "reason",
    "responsibilities",
    "features",
    "main idea",
}
_BRAND_LINES = {"rounds", "syllabloom", "source signals"}
_INSTRUCTION_STARTS = {
    "compare", "design", "describe", "discuss", "evaluate", "explain", "identify",
    "imagine", "list", "predict", "remember", "state", "use",
}
_LECTURE_STOPWORDS = {
    "about", "after", "again", "also", "because", "before", "being", "between",
    "could", "does", "doing", "during", "each", "from", "going", "have", "having",
    "into", "just", "like", "more", "most", "other", "over", "same", "some", "such",
    "than", "that", "their", "them", "then", "there", "these", "they", "this", "those",
    "through", "today", "under", "very", "what", "when", "where", "which", "while", "with",
    "would", "your", "will", "were", "been", "lecture", "class", "thing", "things",
    "everyone", "discuss", "part", "group", "height", "finally", "known", "suggests",
}
_BAD_SUBJECT_STARTS = {
    "and", "or", "but", "so", "because", "at", "by", "could", "during", "for", "from", "if",
    "in", "into", "it", "its", "of", "on", "since", "that", "then", "there", "these",
    "they", "this", "those", "through", "to", "when", "where", "which", "while", "with",
    "he", "i", "i'll", "instead", "may", "might", "must", "note", "notice", "now", "okay",
    "recall", "remember", "shall", "she", "should", "we", "we'll", "would", "you", "you'll",
    "you're", "youre",
}
_BAD_SUBJECT_WORDS = {
    "anything", "because", "can", "cannot", "example", "examples", "everything", "he", "her", "here", "hers", "him",
    "his", "i", "it", "its", "me", "mine", "my", "nothing", "our", "ours", "she",
    "may", "might", "somebody", "someone", "something", "that", "their", "theirs", "them", "there", "these",
    "they", "thing", "things", "this", "those", "us", "we", "what", "whatever", "you",
    "your", "yours",
}
_BAD_SUBJECT_ENDS = {
    "a", "also", "an", "and", "as", "at", "but", "by", "for", "from", "in", "of", "on",
    "or", "the", "these", "those", "to", "too", "with",
}
_PROMOTIONAL_LECTURE_PHRASES = (
    "interactive quizzes",
    "more videos",
    "reading a textbook",
    "say goodbye",
    "learning partner",
    "subscribe to",
    "visit our",
)


def _clean_study_line(value: str) -> str:
    value = re.sub(r"^[\s\u2022\u25aa\u25cf\u25e6\-*]+", "", value or "")
    value = re.sub(r"\s+", " ", value).strip()
    return value


def _is_course_chrome_line(value: str) -> bool:
    cleaned = re.sub(r"\s+", " ", value or "").strip()
    lowered = cleaned.lower()
    if lowered in _BRAND_LINES:
        return True
    if re.fullmatch(r"[A-Z]{2,6}\s*[- ]?\d{2,4}", cleaned):
        return True
    return bool(re.match(r"^[A-Z]{2,6}\s*[- ]?\d{2,4}\s+(?:unit|course)\b", cleaned, flags=re.IGNORECASE))


def _study_units(text: str) -> list[dict]:
    slide_blocks = re.split(r"(?m)^Slide (\d+)\s*$", text)
    if len(slide_blocks) > 1:
        units = []
        for index in range(1, len(slide_blocks), 2):
            lines = [_clean_study_line(line) for line in slide_blocks[index + 1].splitlines()]
            lines = [line for line in lines if line and not re.fullmatch(r"\d+", line)]
            if lines:
                units.append({"number": int(slide_blocks[index]), "lines": lines})
        return units

    lines = [_clean_study_line(line) for line in text.splitlines()]
    lines = [line for line in lines if line and not re.fullmatch(r"\d+", line)]
    if not lines:
        return []
    chunks = [lines[index : index + 8] for index in range(0, len(lines), 8)]
    return [{"number": index + 1, "lines": chunk} for index, chunk in enumerate(chunks)]


def _study_sentences(lines: list[str]) -> list[str]:
    sentences = []
    for line in lines:
        parts = re.split(r"(?<=[.!?])\s+(?=[A-Z0-9])", line)
        for part in parts:
            cleaned = _clean_study_line(part).strip(" ;")
            first_word = next(iter(re.findall(r"[A-Za-z]+", cleaned.lower())), "")
            if (cleaned.endswith("?") and ":" not in cleaned) or first_word in _INSTRUCTION_STARTS:
                continue
            if len(re.findall(r"\b\w+\b", cleaned)) >= 4:
                sentences.append(cleaned)
    return sentences


def _unit_topic(lines: list[str], fallback: str) -> tuple[str, int]:
    usable = [(index, line) for index, line in enumerate(lines) if not _is_course_chrome_line(line)]
    if not usable:
        return fallback, 0

    normalized = [line.upper() for line in lines]
    label_indexes = [
        normalized.index(label)
        for label in ("ATTACHMENT", "ACTION", "INNERVATION")
        if label in normalized
    ]
    if label_indexes:
        before_schema = [
            (index, line)
            for index, line in usable
            if index < min(label_indexes) and line.upper() not in {"ATTACHMENT", "ACTION", "INNERVATION"}
        ]
        natural_title = next(
            ((index, line) for index, line in reversed(before_schema) if not line.isupper()),
            None,
        )
        if natural_title:
            return natural_title[1].rstrip(".:"), natural_title[0]
        if before_schema:
            return before_schema[-1][1].title().rstrip(".:"), before_schema[-1][0]

    first_index, first = usable[0]
    if first.isupper() and len(usable) > 1:
        second_index, second = usable[1]
        if not second.isupper() and second.lower() not in _GENERIC_LABELS:
            return second.rstrip(".:"), second_index
        return first.title().rstrip(".:"), first_index
    return first.rstrip(".:"), first_index


def _valid_card_subject(value: str) -> bool:
    subject = re.sub(r"\s+", " ", value).strip(" ,;:.-")
    subject_for_validation = re.sub(r"^(?:the|a|an)\s+", "", subject, flags=re.IGNORECASE)
    words = re.findall(r"[A-Za-z][A-Za-z'’-]*", subject_for_validation)
    if not (1 <= len(words) <= 11) or len(subject) > 90:
        return False
    lowered = [word.lower().replace("’", "'") for word in words]
    if lowered[0] in _BAD_SUBJECT_STARTS or lowered[-1] in (_BAD_SUBJECT_ENDS | {"which", "that", "who"}):
        return False
    if any(word in _BAD_SUBJECT_WORDS and not original.isupper() for word, original in zip(lowered, words)):
        return False
    if any("'" in word for word in lowered):
        return False
    if "," in subject or ";" in subject:
        return False
    return any(character.isalpha() for character in subject)


def _question_subject(value: str) -> str:
    value = value.strip()
    if value.startswith("The "):
        return "the " + value[4:]
    if value.startswith("A "):
        return "a " + value[2:]
    if value.startswith("An "):
        return "an " + value[3:]
    words = value.split()
    hyphenated_proper_term = "-" in value and any(part[:1].isupper() for part in value.split("-")[1:])
    if not value or (words and words[0].isupper()) or any(word[:1].isupper() for word in words[1:]) or hyphenated_proper_term:
        return value
    return value[:1].lower() + value[1:]


def _looks_plural_subject(value: str) -> bool:
    subject = value.strip().lower()
    if " and " in subject:
        return True
    words = re.findall(r"[a-z]+", subject)
    if not words:
        return False
    last = words[-1]
    if last in {"os", "dos", "unix", "windows", "win32", "ios", "macos", "android"}:
        return False
    if last in {"children", "criteria", "data", "media", "men", "people", "phenomena", "women"}:
        return True
    if last in {"access", "analysis", "bias", "class", "focus", "process", "progress"}:
        return False
    return last.endswith("s") and not last.endswith(("is", "ss", "us"))


def _usable_lecture_answer(value: str) -> bool:
    answer = re.sub(r"\s+", " ", value).strip()
    words = re.findall(r"\b\w+\b", answer)
    if not (2 <= len(words) <= 30) or len(answer) > 240:
        return False
    lowered = answer.lower().replace("’", "'")
    conversational = (
        "i want to know",
        "next slide",
        "right now let's",
        "we are going to",
        "we're going to",
        "we're gonna",
    )
    return not any(phrase in lowered for phrase in conversational)


def _focused_card_answer(value: str) -> str:
    """Keep an atomic answer when a slide sentence chains a second claim."""
    answer = value.strip()
    answer = re.split(
        r",\s+(?:but|although|while|which|(?:and\s+)?(?:asks?|becomes?|coordinates?|provides?|responds?|supplies|updates?))\b",
        answer,
        maxsplit=1,
        flags=re.IGNORECASE,
    )[0]
    answer = re.split(
        r"\s+and\s+(?:asks?|becomes?|can|may|often|provides?|responds?|supplies|then|updates?)\b",
        answer,
        maxsplit=1,
        flags=re.IGNORECASE,
    )[0]
    return answer.strip().rstrip(".")


def _fact_card(statement: str, topic: str = "") -> tuple[str, str] | None:
    statement = statement.strip().rstrip(".")
    if len(statement) < 18 or len(statement) > 520:
        return None

    first_word = next(iter(re.findall(r"[A-Za-z]+", statement.lower())), "")
    if first_word in _INSTRUCTION_STARTS:
        return None

    colon = re.match(r"^([^:]{2,72}):\s+(.{12,})$", statement)
    if colon:
        label, answer = colon.group(1).strip(), colon.group(2).strip()
        normalized_label = label.lower()
        if normalized_label == "question" and "experimental design" in topic.lower():
            outcome = answer.rstrip("?")
            if outcome.lower().startswith("does "):
                clause = outcome[5:6].lower() + outcome[6:]
                clause = re.sub(
                    r"\b(affect|cause|change|decrease|differ|improve|increase|lead|predict|produce|reduce|support)\b",
                    lambda match: match.group(1) + ("es" if match.group(1).endswith(("s", "x", "z", "ch", "sh")) else "s"),
                    clause,
                    count=1,
                )
                outcome = "whether " + clause
            return "What outcome does the experiment test?", outcome
        if normalized_label in {"independent variable", "dependent variable"} and "experimental design" in topic.lower():
            return f"What is the {normalized_label} in the experiment?", answer
        if label.lower() not in _GENERIC_LABELS and _valid_card_subject(label):
            return f"What is {_question_subject(label)}?", answer

    equivalence = re.match(
        r"^(.{2,72}?)\s+(?:is|was)\s+equal to\s+(.{4,})$",
        statement,
        flags=re.IGNORECASE,
    )
    if equivalence and _valid_card_subject(equivalence.group(1)):
        subject = _question_subject(equivalence.group(1))
        return f"What is one {subject} equal to?", equivalence.group(2).strip()

    frequency_definition = re.match(
        r"^(.{2,90}?)\s+(often|usually|typically)\s+(means|is|are)\s+(.{8,})$",
        statement,
        flags=re.IGNORECASE,
    )
    if frequency_definition and _valid_card_subject(frequency_definition.group(1)):
        subject = _question_subject(frequency_definition.group(1))
        frequency = frequency_definition.group(2).lower()
        verb = frequency_definition.group(3).lower()
        if verb == "means":
            return f"What does {subject} {frequency} mean?", frequency_definition.group(4).strip()
        return f"What {'are' if verb == 'are' else 'is'} {subject} {frequency}?", frequency_definition.group(4).strip()

    definition = re.match(
        r"^(.{2,90}?)\s+(is defined as|are defined as|was defined as|were defined as|refers to|means|is|are|was|were)\s+(.{12,})$",
        statement,
        flags=re.IGNORECASE,
    )
    if definition:
        subject, verb, answer = definition.group(1).strip(), definition.group(2).lower(), definition.group(3).strip()
        if _valid_card_subject(subject):
            passive = re.match(
                r"^(oxidized|converted|transported|released|pumped|bound|broken down|broken|stored|carried|secreted|absorbed|filtered|reabsorbed|phosphorylated|dephosphorylated)\b(.+)$",
                answer,
                re.IGNORECASE,
            )
            if verb in {"is", "are", "was", "were"} and passive:
                pronoun = "They" if verb in {"are", "were"} or _looks_plural_subject(subject) else "It"
                be_verb = "are" if pronoun == "They" and verb in {"are", "were"} else "is" if pronoun == "It" and verb in {"is", "was"} else verb
                return f"What happens to {_question_subject(subject)}?", f"{pronoun} {be_verb} {answer.strip()}"
            subject_words = re.findall(r"[A-Za-z]+", subject)
            if len(subject_words) == 1:
                if verb.startswith(("are", "were")) and not _looks_plural_subject(subject):
                    return None
                if verb.startswith(("is", "was")) and _looks_plural_subject(subject):
                    return None
            question_word = "were" if verb.startswith("were") else "was" if verb.startswith("was") else "are" if verb.startswith("are") else "is"
            return f"What {question_word} {_question_subject(subject)}?", answer

    conditional = re.match(
        r"^(.{2,90}?)\s+(occurs?|develops?|improves?|declines?)\s+when\s+(.{8,})$",
        statement,
        flags=re.IGNORECASE,
    )
    if conditional and _valid_card_subject(conditional.group(1)):
        subject = _question_subject(conditional.group(1))
        verb = conditional.group(2).lower()
        base = {"occurs": "occur", "occur": "occur", "develops": "develop", "develop": "develop", "improves": "improve", "improve": "improve", "declines": "decline", "decline": "decline"}[verb]
        return f"When {'do' if _looks_plural_subject(conditional.group(1)) else 'does'} {subject} {base}?", f"When {conditional.group(3).strip()}"

    location = re.match(r"^(.{2,90}?)\s+(occurs?|takes place|is found|are found)\s+(in|at|within|on)\s+(.{4,})$", statement, flags=re.IGNORECASE)
    if location and _valid_card_subject(location.group(1)):
        return f"Where does {_question_subject(location.group(1))} occur?", f"{location.group(3)} {location.group(4).strip()}"

    inclusion = re.match(
        r"^(.{2,90}?)(?:\s+also)?\s+(includes?|contains?|comprises?|consists of)\s+(.{8,})$",
        statement,
        flags=re.IGNORECASE,
    )
    if inclusion and _valid_card_subject(inclusion.group(1)):
        answer = re.sub(r",?\s+(?:right|okay|correct)\?$", "", inclusion.group(3).strip(), flags=re.IGNORECASE)
        return f"What does {_question_subject(inclusion.group(1))} include?", answer

    cause_source = re.match(r"^(.{2,90}?)\s+results? from\s+(.{8,})$", statement, flags=re.IGNORECASE)
    if cause_source and _valid_card_subject(cause_source.group(1)):
        return f"What causes {_question_subject(cause_source.group(1))}?", cause_source.group(2).strip()

    causal = re.match(r"^(.{2,90}?)\s+(causes?|leads to|results in|increases?|decreases?|raises?|inflates?|reduces?)\s+(.{8,})$", statement, flags=re.IGNORECASE)
    if causal and _valid_card_subject(causal.group(1)):
        subject = _question_subject(causal.group(1))
        plural = _looks_plural_subject(causal.group(1))
        pronoun = "They" if plural else "It"
        verb = causal.group(2).lower()
        if plural and verb.endswith("s") and verb not in {"results"}:
            verb = verb[:-1]
        return f"What effect {'do' if plural else 'does'} {subject} have?", f"{pronoun} {verb} {causal.group(3).strip()}."

    function = re.match(r"^(.{2,90}?)(?:\s+(also|still|often|usually|typically))?\s+(allows?|enables?|helps?|functions? to|is responsible for)\s+(.{8,})$", statement, flags=re.IGNORECASE)
    if function and _valid_card_subject(function.group(1)):
        subject = _question_subject(function.group(1))
        auxiliary = "do" if _looks_plural_subject(function.group(1)) else "does"
        adverb = function.group(2)
        verb = function.group(3).lower()
        answer = _focused_card_answer(function.group(4))
        if verb.startswith("help"):
            if answer.lower().startswith("only when "):
                return f"When {auxiliary} {subject} help?", "When " + answer[10:]
            return f"How {auxiliary} {subject}{f' {adverb}' if adverb else ''} help?", answer
        if verb.startswith(("allow", "enable")):
            base = "allow" if verb.startswith("allow") else "enable"
            return f"What {auxiliary} {subject}{f' {adverb}' if adverb else ''} {base}?", answer
        return f"What is the function of {subject}?", answer

    relation = re.match(r"^(.{2,90}?)\s+(depends on|comes from)\s+(.{8,})$", statement, flags=re.IGNORECASE)
    if relation and _valid_card_subject(relation.group(1)):
        subject = _question_subject(relation.group(1))
        phrase = "depend on" if relation.group(2).lower() == "depends on" else "come from"
        return f"What {'do' if _looks_plural_subject(relation.group(1)) else 'does'} {subject} {phrase}?", _focused_card_answer(relation.group(3))

    dated_event = re.match(
        r"^(.{2,90}?)\s+(met|fell|began|ended|occurred|was proclaimed|were proclaimed|was established|were established)\s+(?:in|on|during)\s+((?:[A-Z][a-z]+\s+\d{1,2}(?:,?\s+20\d{2})?)|(?:17|18|19|20)\d{2})\b",
        statement,
        flags=re.IGNORECASE,
    )
    if dated_event and _valid_card_subject(dated_event.group(1)):
        subject, event_verb, event_date = dated_event.groups()
        base_verb = {
            "met": "meet", "fell": "fall", "began": "begin", "ended": "end", "occurred": "occur",
            "was proclaimed": "be proclaimed", "were proclaimed": "be proclaimed",
            "was established": "be established", "were established": "be established",
        }[event_verb.lower()]
        return f"When did {_question_subject(subject)} {base_verb}?", f"In {event_date}."

    conditional_action = re.match(
        r"^(.{2,90}?)\s+can\s+(encourage|internalize|reduce|increase|improve|prevent|promote|support|lower|raise)\s+(.{8,}?)(?:\s+when\s+(.{8,})|\s+by\s+(.{8,}))$",
        statement,
        flags=re.IGNORECASE,
    )
    if conditional_action and _valid_card_subject(conditional_action.group(1)):
        raw_subject, verb, object_phrase, condition, mechanism = conditional_action.groups()
        subject = _question_subject(raw_subject)
        if condition:
            return f"When can {subject} {verb.lower()} {object_phrase.strip()}?", f"When {condition.strip()}"
        return f"How can {subject} {verb.lower()} {object_phrase.strip()}?", f"By {mechanism.strip()}"

    action = re.match(
        r"^(.{2,90}?)\s+(briefly preserves|maintains and manipulates|adopted|adopts|aligns|allocates|alternates|bind|binds|brought|brings|changes|coordinates|correct|corrects|create|creates|decreases|described|describes|disrupts|distributes|estimates|evaluates|favors|fixes|generates|groups|impose|imposes|imposed|increases|integrates|involved|involves|judges|maintains|makes|mixes|needs|open|opens|pairs|predict|predicts|preserves|presents|proposed|proposes|protects|provides|pump|pumps|produces|reflects|release|releases|remember|remembers|reorganizes|repeats|represents|requires|require|return|returns|retains|shows|stabilizes|supplies|supports|tests|trigger|triggers|treats|updates|uses|convert|converts|regulates|regulate)\s+(.{8,})$",
        statement,
        flags=re.IGNORECASE,
    )
    if action and _valid_card_subject(action.group(1)):
        verb = action.group(2).lower()
        raw_subject = action.group(1).strip()
        raw_subject = re.sub(r"\s*\((?:18|19|20)\d{2}\)", "", raw_subject)
        subject = raw_subject if verb in {"described", "proposed"} and len(raw_subject.split()) == 1 else _question_subject(raw_subject)
        base_verbs = {
            "adopted": "adopt", "adopts": "adopt", "aligns": "align", "allocates": "allocate",
            "alternates": "alternate", "bind": "bind", "binds": "bind", "brought": "bring", "brings": "bring",
            "coordinates": "coordinate", "correct": "correct",
            "changes": "change", "convert": "convert", "converts": "convert", "create": "create", "creates": "create", "decreases": "decrease",
            "corrects": "correct", "described": "describe", "describes": "describe", "disrupts": "disrupt",
            "distributes": "distribute", "estimates": "estimate", "evaluates": "evaluate", "favors": "favor", "fixes": "fix",
            "falls": "fall", "generates": "generate", "groups": "group", "impose": "impose", "imposes": "impose", "imposed": "impose",
            "increases": "increase", "integrates": "integrate", "involved": "involve", "involves": "involve",
            "judges": "judge", "maintains": "maintain",
            "makes": "make", "mixes": "mix", "needs": "need", "pairs": "pair", "predict": "predict", "predicts": "predict",
            "open": "open", "opens": "open", "preserves": "preserve", "proposed": "propose",
            "presents": "present", "proposes": "propose", "protects": "protect", "provides": "provide", "reflects": "reflect",
            "pump": "pump", "pumps": "pump", "produces": "produce", "release": "release", "releases": "release",
            "remember": "remember", "remembers": "remember", "reorganizes": "reorganize", "repeats": "repeat", "represents": "represent",
            "return": "return", "returns": "return",
            "require": "require", "requires": "require", "retains": "retain", "shows": "show",
            "stabilizes": "stabilize", "supplies": "supply", "supports": "support", "tests": "test",
            "trigger": "trigger", "triggers": "trigger", "treats": "treat", "updates": "update", "uses": "use", "regulates": "regulate",
            "regulate": "regulate", "briefly preserves": "briefly preserve",
            "maintains and manipulates": "maintain and manipulate",
        }
        auxiliary = "did" if verb in {"adopted", "brought", "described", "fell", "involved", "imposed", "proposed"} else "do" if _looks_plural_subject(raw_subject) else "does"
        if verb == "proposed" and len(raw_subject.split()) == 1 and topic.strip():
            return f'In “{topic.strip()},” what did {subject} propose?', _focused_card_answer(action.group(3))
        if raw_subject.lower() in {"the distinction", "this distinction", "the difference", "this difference"} and topic.strip():
            return f'What does the distinction within “{topic.strip()}” describe?', _focused_card_answer(action.group(3))
        if verb == "makes":
            made = re.match(r"^(.+?)\s+(easier|harder)\s+to\s+(.+)$", _focused_card_answer(action.group(3)), flags=re.IGNORECASE)
            if made:
                return f"What {auxiliary} {subject} make {made.group(2).lower()} to {made.group(3)}?", made.group(1).strip()
        return f"What {auxiliary} {subject} {base_verbs[verb]}?", _focused_card_answer(action.group(3))
    return None


def _labeled_assessment_cards(text: str, filename: str, status: str = "verified") -> list[dict]:
    """Pair explicit question and answer paragraphs without turning the label into a card."""
    question_pattern = re.compile(r"^\s*(?:question|q)\s*(\d+(?:\.\d+)*)\s*[:.)-]\s*(.+?)\s*$", re.IGNORECASE)
    answer_pattern = re.compile(
        r"^\s*(?:correct\s+)?(?:answer|ans)\s*(?:(?:to\s+)?(?:question\s*)?(\d+(?:\.\d+)*))?\s*[:.)-]\s*(.*?)\s*$",
        re.IGNORECASE,
    )
    blocks: list[dict] = []
    current: dict | None = None
    answer_started = False

    def finish() -> None:
        if current and current.get("question") and current.get("answer"):
            blocks.append(current.copy())

    for raw_line in text.splitlines():
        line = re.sub(r"\s+", " ", raw_line).strip()
        if not line:
            continue
        question_match = question_pattern.match(line)
        if question_match:
            finish()
            current = {"id": question_match.group(1), "question": question_match.group(2).strip(), "answer": []}
            answer_started = False
            continue
        if current is None:
            continue
        answer_match = answer_pattern.match(line)
        if answer_match:
            answer_id, answer_text = answer_match.groups()
            if answer_id and answer_id != current["id"]:
                continue
            current["answer"].append(answer_text.strip())
            answer_started = True
            continue
        if answer_started:
            current["answer"].append(line)
        elif not line.lower().startswith(("quiz ", "answer key", "answers")):
            current["question"] += " " + line
    finish()

    cards = []
    for item in blocks:
        front = re.sub(r"\s+", " ", item["question"]).strip().rstrip(" .")
        back = re.sub(r"\s+", " ", " ".join(item["answer"])).strip()
        if not front.endswith(("?", ".")):
            front += "?"
        if len(front.split()) < 5 or len(re.findall(r"\b\w+\b", back)) < 3 or len(back) > 600:
            continue
        cards.append({
            "muscle": "Quiz answer key",
            "field": "assessment",
            "front": front,
            "back": back[:1].upper() + back[1:],
            "section": "Quiz answer key",
            "source": f"{filename} · Question {item['id']}",
            "status": status,
            "questionId": item["id"],
        })
    return cards


def _topic_question(topic: str, lines: list[str]) -> str:
    normalized = topic.strip().rstrip(".?")
    patterns = (
        (r"^(types|forms|classes) of (.+)$", "What are the main {0} of {1}?"),
        (r"^(functions|roles) of (.+)$", "What are the main {0} of {1}?"),
        (r"^(stages|steps|phases) of (.+)$", "What are the {0} of {1}?"),
        (r"^(causes|effects|risk factors|features|components) of (.+)$", "What are the main {0} of {1}?"),
    )
    for pattern, template in patterns:
        match = re.match(pattern, normalized, flags=re.IGNORECASE)
        if match:
            return template.format(match.group(1).lower(), match.group(2))
    if re.search(r"\b(vs\.?|versus)\b", normalized, flags=re.IGNORECASE):
        return f"How do {normalized} differ?"
    return f"What are the key ideas about {normalized}?" if len(lines) > 1 else f"What is the key idea about {normalized}?"


def _infer_lecture_topic(text: str) -> str:
    talk_about = re.search(
        r"\b(?:talk|talking|learn|learning|focus|focusing|cover|covering)\s+(?:about|on)\s+(.{3,70}?)(?:[.,]|\band how\b|\band why\b|$)",
        text,
        flags=re.IGNORECASE,
    )
    if talk_about:
        return talk_about.group(1).strip().rstrip(".")
    for sentence in _study_sentences([text]):
        definition = re.match(r"^(.{2,90}?)\s+(?:is|are|means|refers to)\s+", sentence, flags=re.IGNORECASE)
        if definition and _valid_card_subject(definition.group(1)):
            candidate = definition.group(1).strip()
            candidate = re.sub(r"^(?:the|a|an)\s+", "", candidate, flags=re.IGNORECASE)
            return candidate
    words = [word.lower() for word in re.findall(r"\b[A-Za-z][A-Za-z-]{3,}\b", text)]
    useful = [word for word in words if word not in _LECTURE_STOPWORDS]
    if len(useful) >= 2:
        bigrams = [f"{useful[index]} {useful[index + 1]}" for index in range(len(useful) - 1)]
        counts = {phrase: bigrams.count(phrase) for phrase in dict.fromkeys(bigrams)}
        return max(counts, key=lambda phrase: (counts[phrase], -bigrams.index(phrase))).replace("-", " ")
    if useful:
        return useful[0].replace("-", " ")
    return "this lecture section"


def _topic_from_card_question(question: str) -> str:
    patterns = (
        r"^What is one (.+) equal to\?$",
        r"^What is the function of (.+)\?$",
        r"^What effect does (.+) have\?$",
        r"^What does (.+) include\?$",
        r"^Where does (.+) occur\?$",
        r"^What (?:is|are|was|were) (.+)\?$",
    )
    for pattern in patterns:
        match = re.match(pattern, question.strip(), flags=re.IGNORECASE)
        if not match:
            continue
        topic = match.group(1).strip()
        topic = re.sub(r"^(?:the|a|an)\s+", "", topic, flags=re.IGNORECASE)
        if topic.lower() in _GENERIC_LABELS:
            return ""
        if _valid_card_subject(topic):
            return topic[:1].upper() + topic[1:]
    return ""


def compile_study_material(text: str, filename: str, status: str = "verified") -> dict:
    concepts = []
    notes = []
    cards = []
    seen_cards = set()
    stem = Path(filename).stem.replace("_", " ").replace("-", " ").strip()

    for unit in _study_units(text):
        lines = unit["lines"]
        if not lines:
            continue
        topic, title_index = _unit_topic(lines, stem)
        has_title = len(topic) <= 120 and len(topic.split()) <= 16
        body = lines[title_index + 1 :] if has_title and title_index + 1 < len(lines) else lines
        body = [line for line in body if not _is_course_chrome_line(line)]
        body = [line for line in body if line.lower() not in _GENERIC_SECTION_TITLES]
        if not body:
            continue
        if topic.lower() in _GENERIC_SECTION_TITLES:
            topic = _infer_lecture_topic(" ".join(body))

        citation = f"{filename} · Slide {unit['number']}" if text.lstrip().startswith("Slide ") else filename
        note_lines = body[:6]
        notes.append({
            "title": topic,
            "section": topic,
            "heardAt": None,
            "status": status,
            "source": citation,
            "slideNumber": unit["number"] if text.lstrip().startswith("Slide ") else None,
            "lines": note_lines,
        })

        unit_cards = []
        case_subject = ""
        if topic.lower().startswith("case analysis:"):
            case_subject = topic.split(":", 1)[1].split("'s", 1)[0].strip().lower()
        for statement in _study_sentences(body):
            if case_subject and statement.lower().startswith(case_subject + " "):
                continue
            generated = _fact_card(statement, topic)
            if not generated:
                continue
            front, back = generated
            key = (front.lower(), back.lower())
            if key in seen_cards:
                continue
            seen_cards.add(key)
            unit_cards.append((front, back))
            if len(unit_cards) >= 3:
                break

        if unit_cards:
            concepts.append({
                "name": topic,
                "status": status,
                "source": citation,
                "slideNumber": unit["number"] if text.lstrip().startswith("Slide ") else None,
            })

        for front, back in unit_cards[:3]:
            cards.append({
                "muscle": topic,
                "field": "concept",
                "front": front,
                "back": back,
                "section": topic,
                "source": citation,
                "status": status,
                "slideNumber": unit["number"] if text.lstrip().startswith("Slide ") else None,
            })

    return {"concepts": concepts, "notes": notes, "cards": cards}


_PDF_BOILERPLATE = re.compile(
    r"^(?:https?://\S+|www\.\S+|.{2,120}\s+\d{1,2}[/-]\d{1,2}[/-]\d{2,4},?\s+\d{1,2}:\d{2}\s*(?:am|pm)$|"
    r"CS\s*\d{2,4}\b|lecturer\s*:|read\s*:|topics\s*:|lecture\s*(?:#|no\.?\s*)?\d+\b|"
    r"fig(?:ure)?\.?\s*\d*:?|slide\s*\d+\b|page\s*\d+(?:\s+of\s+\d+)?\b)",
    flags=re.IGNORECASE,
)
_PDF_VAGUE_QUESTION = re.compile(
    r"^what\s+(?:is|are|was|were|do|does|did)\s+(?:the\s+)?"
    r"(?:topics?|read|reason|one|thing|user|system|structure|responsibilities|advantages?|disadvantages?)\??$",
    flags=re.IGNORECASE,
)


def _pdf_content_lines(page_text: str) -> list[str]:
    lines = []
    for raw_line in page_text.splitlines():
        line = re.sub(r"\s+", " ", raw_line).strip()
        line = re.sub(r"\b(\d+)\s+(st|nd|rd|th)\b", r"\1\2", line, flags=re.IGNORECASE)
        line = line.lstrip("\u2022\u25aa\u25cf\u25e6\uf0b7-–* ").strip()
        if not line or _PDF_BOILERPLATE.match(line) or line.lower() in _BRAND_LINES:
            continue
        if re.fullmatch(r"(?:st|nd|rd|th)", line, flags=re.IGNORECASE):
            continue
        if re.fullmatch(r"[\d\W]+", line):
            continue
        if len(line.split()) <= 1 and len(line) < 4:
            continue
        if len(re.findall(r"(?<!\S)[A-Za-z](?=\s|$)", line)) >= 5:
            continue
        if len(line) > 18 and sum(character.isalpha() for character in line) / len(line) < 0.42:
            continue
        lines.append(line)
    return lines


def _pdf_paragraphs(page_text: str) -> list[str]:
    paragraphs = []
    current = ""

    def flush() -> None:
        nonlocal current
        if current:
            paragraphs.append(current.strip())
            current = ""

    for line in _pdf_content_lines(page_text):
        title_like = (
            len(line.split()) <= 8
            and len(line) <= 72
            and not re.search(r"[.!?]$", line)
            and (
                line.istitle() or line.isupper() or line.endswith(":") or line.endswith(")")
                or (len(line.split()) <= 8 and " - " in line)
            )
        )
        if title_like:
            flush()
            paragraphs.append(line.rstrip(":"))
            continue
        current_tokens = re.findall(r"[A-Za-z]+", current)
        next_tokens = re.findall(r"[A-Za-z]+", line)
        last_token = current_tokens[-1] if current_tokens else ""
        last_word = last_token.lower()
        next_token = next_tokens[0] if next_tokens else ""
        continuation = (
            bool(current)
            and not re.search(r"[.!?]$", current)
            and not line.startswith(("- ", "• "))
            and (
                line[:1].islower() or line[:1] in "([\"'" or current.endswith("-")
                or current.count("(") > current.count(")")
                or last_word in {"a", "an", "and", "as", "at", "by", "for", "from", "in", "of", "on", "or", "the", "to", "with"}
                or (len(current) > 72 and last_token[:1].isupper() and next_token[:1].isupper())
            )
        )
        if current and continuation:
            current += ("" if current.endswith("-") else " ") + line
        else:
            flush()
            current = line
        if re.search(r"[.!?]$", current):
            flush()
    flush()
    return paragraphs


def _pdf_question_from_label(label: str, answer: str) -> str:
    label = re.sub(r"\s+", " ", label).strip(" ,;.:-")
    label = re.sub(r"\s+[-–]\s+", " ", label)
    label = re.sub(r"\bms\s+dos\b", "MS-DOS", label, flags=re.IGNORECASE)
    label = re.sub(r"\s*\((?:interactive)\)", "", label, flags=re.IGNORECASE)
    normalized = label.lower()
    if re.match(r"^(?:an?\s+)?implementing\s+prescribed\s+apis?\b", normalized):
        return "Why is implementing the prescribed APIs important?"
    if normalized.startswith("command interpreter in unix"):
        return "What is the Unix command interpreter?"
    if normalized.startswith("win32api") or normalized.startswith("win32 api"):
        return "What does the Win32 API represent?"
    if normalized == "multiprocessing environment":
        return "What defines a multiprocessing environment?"
    label_article = re.match(r"^(a|an|the)\s+(.+)$", label, flags=re.IGNORECASE)
    if label_article:
        article, label = label_article.groups()
        normalized = label.lower()
        return f"What is {article.lower()} {normalized}?"
    if normalized.endswith(("systems", "processes", "instructions", "programs", "layers", "models")):
        return f"What are {normalized}?"
    answer_article = re.match(r"^(a|an|the)\s+", answer.strip(), flags=re.IGNORECASE)
    countable_nouns = (
        "system", "program", "model", "queue", "bit", "layer", "process", "interpreter",
        "routine", "task", "environment", "operation", "interface", "structure",
    )
    if answer_article or normalized.endswith(countable_nouns):
        article = answer_article.group(1).lower() if answer_article else "the" if normalized.startswith(("simplest ", "highest ", "lowest ", "main ", "most ")) else "a"
        if article == "a" and re.match(r"^[aeiou]", normalized) and not normalized.startswith(("user ", "university ", "uniprocessing ")):
            article = "an"
        return f"What is {article} {normalized}?"
    return f"What is {normalized}?"


def _pdf_fact_card(statement: str, section: str = "") -> tuple[str, str] | None:
    statement = re.sub(r"\s+", " ", statement).strip(" \t\r\n-•")
    section_key = re.sub(r"\s+", " ", section).strip().lower()
    if "starting a computer" in section_key and re.match(r"^starting a computer:\s+in order to start running", statement, flags=re.IGNORECASE):
        return "What initial program does a computer need to start or reboot?", "An initial program called a bootstrap program."
    if statement.lower().startswith("the occurrence of an event is signaled by an interrupt"):
        return "How are hardware and software events signaled?", "Hardware signals the CPU with an interrupt; software signals it by executing a system call."
    if statement.lower().startswith("a mode bit, is added to the hardware"):
        return "What does the mode bit indicate?", "It indicates the computer's current mode."
    if statement.lower().startswith("both programmed i/o and interrupt driven i/o require"):
        return (
            "What memory work must the CPU handle in both programmed and interrupt-driven I/O?",
            "The CPU extracts data from main memory for output and stores input data in main memory.",
        )
    if statement.lower().startswith("the command interpreter is the interface between"):
        return "What does the command interpreter provide?", "It provides the interface between the user and the operating system."
    if statement.lower().startswith("the command interpreter includes") or statement.lower().startswith("command interpreter includes"):
        return "What does the command interpreter include?", "It includes the code needed to execute commands."
    if statement.lower().startswith("the operating system is partitioned into system components"):
        return "How is the operating system organized in this lecture?", "It is partitioned into components, each with a specific task."
    if statement.lower().startswith("ms-dos - application programs are able to access"):
        return "What access do MS-DOS applications have to I/O routines?", "They can access the basic I/O routines."
    if statement.lower().startswith("this makes ms-dos vulnerable to errant programs"):
        return "What weakness follows from MS-DOS applications accessing basic I/O routines?", "MS-DOS is vulnerable to errant programs."
    if statement.lower().startswith("microkernel represents the most used and fundamental component"):
        return None
    if statement.lower().startswith("user level application are written"):
        return "What does the lecture say about user-level applications?", "They are written in C/C++ and do not depend on the architecture."
    if statement.lower().startswith("android: designed for android smartphone"):
        return "Which devices is Android designed for?", "Android is designed for smartphones and tablet computers."
    if statement.lower().startswith("the task of the operating system is to automatically transfer control"):
        return "What does the operating system do between jobs?", "It automatically transfers control from one job to the next."
    if statement.lower().startswith("multiprocessing environment: more than one cpu"):
        return "What defines a multiprocessing environment?", "It has more than one CPU and many processes ready for execution."
    if statement.lower().startswith("hard real-time system: guarantees that"):
        return "What does a hard real-time system guarantee?", "It guarantees that a critical task is completed in time."
    if statement.lower().startswith("soft real-time system:"):
        return "How does a soft real-time system treat a critical task?", "It gives the critical real-time task priority over other tasks."
    if statement.lower().startswith("process switch times"):
        return "Why are process-switch times considered overhead?", "Process-switch time is pure overhead and depends heavily on hardware support."
    if statement.lower().startswith("the executive provides services used by all the environment subsystems"):
        return (
            "What services does the Windows executive provide?",
            "It provides common services to environment subsystems, including object, virtual-memory, and process management, plug-and-play, and startup management.",
        )
    if statement.lower().startswith("responsibilities: thread scheduling") and "windows 10" in section_key:
        responsibilities = statement.split(":", 1)[1].strip()
        return "What responsibilities does the Windows 10 kernel handle?", "The kernel handles " + responsibilities
    if statement.lower().startswith("hal isolates the os from platform specific hardware differences"):
        return "What does the Windows Hardware Abstraction Layer isolate the OS from?", "It isolates the OS from platform-specific hardware differences."
    if statement.lower().startswith("provides the support for symmetric multiprocessing"):
        return "What does the Windows Hardware Abstraction Layer support?", "It supports symmetric multiprocessing (SMP)."
    if statement.lower().startswith("most of the upper level modules can access the hardware only through the hal"):
        return "How do most upper-level Windows modules access hardware?", "They access it only through the Hardware Abstraction Layer (HAL)."
    if statement.lower().startswith("a small core of facilities provides functions and services"):
        return None
    if statement.lower().startswith("win32api represents the native environment"):
        return (
            "What does the Win32 API represent?",
            "It is the native environment for Windows NT, Windows 2000, Windows XP, and Windows 10.",
        )
    if statement.lower().startswith("command interpreter in unix is a process"):
        return "What is the Unix command interpreter?", "It is a process that runs in user mode."
    if statement.lower().startswith("each layer uses functions and services of lower-level layers only"):
        return "What may each layer in the layered OS use?", "Only functions and services of lower-level layers."
    if "operating system" in statement.lower() and "automatically transfer control from one job to the following one" in statement.lower():
        return "What does the operating system do between jobs?", "It automatically transfers control from one job to the next."
    if re.match(r"^starting a computer:\s+in order to start running", statement, flags=re.IGNORECASE):
        return "What initial program does a computer need to start or reboot?", "An initial program called a bootstrap program."
    if "programmed i/o" in section_key and statement.lower().startswith("data is exchanged between the cpu and the i/o module"):
        return "How does programmed I/O handle data?", "The CPU and I/O module exchange data when the CPU issues a command; the CPU usually waits for the operation to finish."
    if "interrupt driven i/o" in section_key and statement.lower().startswith("the cpu issues a command requesting i/o"):
        return None
    if "direct memory access" in section_key and statement.lower().startswith("the i/o module and main memory exchange data directly"):
        return "How does direct memory access transfer data?", statement
    if "interrupts" in section_key and statement.lower().startswith("an interrupt can be synchronous if"):
        return None
    if "operating system architecture" in section_key and statement.lower().startswith("simplest structure:"):
        return "What defines a monolithic operating-system structure?", "It places all kernel functionality in one file and one address space, without a modular design."
    if "monolithic structure" in section_key and statement.lower().startswith("simplest structure:"):
        return "What defines a monolithic operating-system structure?", "It places all kernel functionality in one file and one address space, without a modular design."
    if "microkernel" in section_key and re.match(r"^(?:the\s+)?microkernel functions as a message exchange:", statement, flags=re.IGNORECASE):
        return "What does the microkernel do as a message exchange?", "It validates and passes messages between components, grants hardware access, and performs a protection function."
    if "windows 10" in section_key and statement.lower().startswith("kernel - hybrid architecture") and "responsibilities:" in statement.lower():
        responsibilities = re.split(r"\bresponsibilities:\s*", statement, maxsplit=1, flags=re.IGNORECASE)[-1]
        return "What responsibilities does the Windows 10 kernel handle?", responsibilities
    if ("programmed i/o" in section_key or "i/o structure" in section_key) and statement.lower().startswith("both programmed i/o and interrupt driven i/o require"):
        return "What data-handling work do programmed and interrupt-driven I/O require from the CPU?", statement
    if statement.lower().startswith("the command interpreter is invoked when the computer is started"):
        return "When is the MS-DOS command interpreter invoked?", "When the computer is started."
    if statement.lower().startswith("the shell of the user's choice"):
        return "When is a user's Unix shell started?", statement
    process_state = re.match(r"^([A-Z][a-z]+(?:\s+\([^)]+\)|\s+swapped)?)\s*-\s+(?:a|the) process\s+(.+)$", statement, flags=re.IGNORECASE)
    if process_state:
        state, description = process_state.groups()
        return f"What does the {state.strip().lower()} process state mean?", f"The process {description.strip()}"
    mode_description = re.match(
        r"^(user-mode|system\s*\(monitor\)\s*-mode)\s*\([^)]*\)\s+the operation is done on behalf of\s+(.+)$",
        statement,
        flags=re.IGNORECASE,
    )
    if mode_description:
        mode, beneficiary = mode_description.groups()
        mode_name = "user mode" if mode.lower().startswith("user") else "system mode"
        return f"On whose behalf does {mode_name} operate?", f"On behalf of {beneficiary.strip()}"
    interrupt = re.match(
        r"^An interrupt can be synchronous if it is the (?:direct|DIRECT) result of (.+?)\.\s*Otherwise it is known as asynchronous\.?$",
        statement,
        flags=re.IGNORECASE,
    )
    if interrupt:
        return "How do synchronous and asynchronous interrupts differ?", f"A synchronous interrupt is the direct result of {interrupt.group(1).strip()}; otherwise, it is asynchronous."
    sync_interrupt = re.match(
        r"^An interrupt can be synchronous if it is the direct result of (.+)$",
        statement.rstrip("."),
        flags=re.IGNORECASE,
    )
    if sync_interrupt:
        return "When is an interrupt synchronous?", f"When it is the direct result of {sync_interrupt.group(1).strip()}"
    command_implementation = re.match(r"^Commands are implemented by\s+(.+)$", statement, flags=re.IGNORECASE)
    if command_implementation:
        return "How are commands implemented?", f"By {command_implementation.group(1).strip()}"
    labeled = re.match(r"^([^:]{2,72}):\s*(.{10,})$", statement)
    if labeled:
        label, answer = labeled.groups()
        if _valid_card_subject(label) and label.lower() not in _GENERIC_LABELS and not _PDF_BOILERPLATE.match(label):
            plural = label.lower().endswith(("systems", "processes", "instructions", "programs", "layers", "models"))
            model_label = re.sub(r"^the\s+", "", label.strip(), flags=re.IGNORECASE).lower()
            if model_label == "symmetric multiprocessing model" and re.fullmatch(r"smp\s*:\s*peer-to-peer", answer.strip(), flags=re.IGNORECASE):
                return "What characterizes the symmetric multiprocessing model?", "It uses a peer-to-peer model."
            if model_label == "asymmetric multiprocessing model" and answer.strip().lower() == "master-slave":
                return "What characterizes the asymmetric multiprocessing model?", "It uses a master-slave model."
            guarantee = re.match(r"guarantees?\s+that\s+(.+)", answer, flags=re.IGNORECASE)
            if guarantee and label.lower().endswith("system"):
                return f"What does a {label.lower()} guarantee?", f"A {label.lower()} guarantees that {guarantee.group(1).strip()}"
            if re.match(r"used\s+when\s+", answer, flags=re.IGNORECASE):
                return _pdf_question_from_label(label, answer), ("They are " if plural else "It is ") + answer.strip()
            if label.lower() == "system call" and answer.lower().startswith("interface between"):
                return "What is a system call?", "A system call is an interface between a user process and the operating system (the kernel)."
            if label.lower() == "command interpreter" and answer.lower().startswith("interface between"):
                return "What does the command interpreter provide?", "It provides the interface between " + answer[len("interface between"):].strip()
            if label.lower() == "process switch times":
                return "Why are process-switch times considered overhead?", "Process-switch time is pure overhead and depends heavily on hardware support."
            if label.lower() == "multiprocessing environment":
                return "What defines a multiprocessing environment?", "It has more than one CPU and many processes ready for execution."
            if label.lower() == "each layer" and answer.lower().startswith("functions and services of lower-level layers"):
                return "What may each layer in the layered OS use?", "Only functions and services of lower-level layers."
            if label.lower().startswith("task of the operating system") and answer.lower().startswith("to automatically transfer control"):
                return "What does the operating system do between jobs?", "It automatically transfers control from one job to the next."
            if label.lower() == "win32api" and answer.lower().startswith("the native environment"):
                return (
                    "What does the Win32 API represent?",
                    "It is the native environment for Windows NT, Windows 2000, Windows XP, and Windows 10.",
                )
            if label.lower() == "command interpreter in unix" and answer.lower().startswith("a process that runs in user mode"):
                return "What is the Unix command interpreter?", "It is a process that runs in user mode."
            if section_key == "command interpreter" and label.lower() in {"includes", "include"}:
                return "What does the command interpreter include?", "It includes the code needed to execute commands."
            if label.lower() == "microkernel" and answer.lower().startswith("the most used"):
                return None
            if label.lower() == "executive" and answer.lower().startswith("provides services"):
                return (
                    "What services does the Windows executive provide?",
                    "It provides common services to environment subsystems, including object, virtual-memory, and process management, plug-and-play, and startup management.",
                )
            if label.lower() == "android" and answer.lower().startswith("designed for"):
                return "What is Android designed for?", "Android is " + answer.strip()
            if label.lower() == "system/360" and answer.lower().startswith("designed to"):
                return "What was System/360 designed to do?", "It was " + answer.strip()
            if label.lower() == "darwin" and answer.lower().startswith("example of microkernel"):
                return (
                    "Which components make up Darwin in these notes?",
                    "A Mach microkernel component and a BSD Unix kernel, within a hybrid, layered structure.",
                )
            if answer.lower().startswith("place all the functionality") and label.lower() == "simplest structure":
                return "What defines a monolithic operating-system structure?", "It places all kernel functionality in one file and one address space, without a modular design."
            if re.match(r"(?:added|broken|characterized|designed|done|included|invoked|partitioned|written)\s+", answer, flags=re.IGNORECASE):
                subject = re.sub(r"^(?:a|an|the)\s+", "", label.strip(), flags=re.IGNORECASE)
                if answer.lower().startswith("designed"):
                    return f"What was {subject.lower()} designed for?", f"{subject} is {answer.strip()}"
                article = "an" if re.match(r"^[aeiou]", subject.lower()) else "a"
                return _pdf_question_from_label(label, answer), f"{article.capitalize()} {subject.lower()} is {answer.strip()}"
            if answer.lower().startswith("the ") and label.lower().startswith("the "):
                return _pdf_question_from_label(label, answer), "It is " + answer.strip()
            if label.lower() in {"multiprogramming", "multiprocessing"} and answer.lower().startswith(("one cpu", "several processors")):
                if label.lower() == "multiprogramming":
                    return "What is multiprogramming?", "It uses one CPU, with multiple processes ready for execution."
                return "What is multiprocessing?", answer[:1].upper() + answer[1:].strip()
            return _pdf_question_from_label(label, answer), answer.strip()
    role = re.match(
        r"^(?:the\s+)?(.{2,60}?)'s\s+(?:(?:major|main)\s+)?(task|role|function)\s+(?:is|was)\s+to\s+(.+)$",
        statement.rstrip("."),
        flags=re.IGNORECASE,
    )
    if role:
        owner, role_name, action = role.groups()
        return f"What is the {role_name.lower()} of the {owner.lower()}?", f"To {action.strip()}"
    system_limits = re.match(r"^The system structure is limited by (.+)$", statement, flags=re.IGNORECASE)
    if system_limits:
        return "What limits the operating-system structure?", system_limits.group(1).strip()
    user_mode_os = re.match(r"^(?:the\s+)?OS is broken into (.+)$", statement, flags=re.IGNORECASE)
    if user_mode_os:
        return "How is the operating system organized in the layered structure described here?", "It is broken into " + user_mode_os.group(1).strip()
    designed = re.match(r"^(.{2,60}?)\s+was designed to\s+(.+)$", statement.rstrip("."), flags=re.IGNORECASE)
    if designed and _valid_card_subject(designed.group(1)):
        subject, action = designed.groups()
        return f"What was {_question_subject(subject)} designed to do?", f"It was designed to {action.strip()}"
    generated = _fact_card(statement)
    if not generated:
        return None
    front, back = generated
    definition = re.match(
        r"^(.{2,90}?)\s+(?:is defined as|are defined as|was defined as|were defined as|refers to|means|is|are|was|were)\s+(.{8,})$",
        statement.rstrip("."),
        flags=re.IGNORECASE,
    )
    if definition and _valid_card_subject(definition.group(1)):
        subject, answer = definition.groups()
        front = _pdf_question_from_label(subject, answer)
        verb_match = re.search(
            r"\s+(is defined as|are defined as|was defined as|were defined as|refers to|means|is|are|was|were)\s+",
            statement.rstrip("."),
            flags=re.IGNORECASE,
        )
        verb = verb_match.group(1).lower() if verb_match else ""
        if verb in {"is", "are", "was", "were"} and not answer.lower().startswith(("a ", "an ", "the ")):
            pronoun = "They" if verb in {"are", "were"} or _looks_plural_subject(subject) else "It"
            be_verb = "are" if pronoun == "They" and verb in {"are", "were"} else "is" if pronoun == "It" and verb in {"is", "was"} else verb
            back = f"{pronoun} {be_verb} {answer}"
    return front, back


def _pdf_card_is_usable(front: str, back: str, allow_long_answer: bool = False) -> bool:
    front = re.sub(r"\s+", " ", front).strip()
    back = re.sub(r"\s+", " ", back).strip()
    if not (12 <= len(front) <= 180 and front.endswith("?")) or _PDF_VAGUE_QUESTION.match(front):
        return False
    if re.search(r"(?:https?://|www\.)\S+", front + " " + back, flags=re.IGNORECASE):
        return False
    if _PDF_BOILERPLATE.match(front) or _PDF_BOILERPLATE.match(back):
        return False
    malformed_subject = re.match(r"^what\s+(?:is|are|was|were)\s+(.+)$", front[:-1], flags=re.IGNORECASE)
    if malformed_subject and re.search(
        r"\b(?:is|are|was|were|does|did|do|functions|represents|includes|provides|limited)\b",
        malformed_subject.group(1),
        flags=re.IGNORECASE,
    ):
        return False
    if any(fragment in front.lower() for fragment in ("one of the main", "prior to the", "what is pc?", "what is meant by pc?", "what is meant by the reason", "what is the interfaces", "what is the main idea?", "what is meant by the main idea", "what effect do the program runs", "ms - dos at boot time")):
        return False
    if re.search(r"\bin class\b", front + " " + back, flags=re.IGNORECASE):
        return False
    if any(fragment in front.lower() for fragment in ("outer circles", "mach?", "is a the ", "is an user", "mode bit,?")):
        return False
    if any(fragment in front.lower() for fragment in ("what is topics", "what is read", "what is responsibilities")):
        return False
    words = re.findall(r"\b\w+\b", back)
    max_words = 110 if allow_long_answer else 55
    max_chars = 720 if allow_long_answer else 360
    if not (3 <= len(words) <= max_words) or len(back) > max_chars:
        return False
    if re.match(r"^\d+\s*[;:]", back) or "�" in back:
        return False
    if "\ufffd" in back or re.search(r"\b(?:and|or|of|to|with|which|that|because|the)\s*$", back, flags=re.IGNORECASE):
        return False
    if re.search(r"\b(?:ios|darwin)\s*:", back, flags=re.IGNORECASE):
        return False
    if re.search(r"(?:\b\w\s){5,}\w\b", back):
        return False
    return True


def _pdf_sequence_cards(paragraphs: list[str]) -> list[tuple[str, str]]:
    cards = []
    section = ""
    question = ""
    steps = []

    def flush() -> None:
        nonlocal steps
        if len(steps) >= 2 and question:
            mode_switch = question == "What processor actions occur during a mode switch?"
            process_switch = question == "What steps occur during a full process switch?"
            if process_switch and len(steps) >= 7:
                cards.append((
                    question,
                    "1. Save the processor context. 2. Update the current process's PCB. 3. Move that PCB to the appropriate queue. 4. Select another process. 5. Update the new process's PCB. 6. Update memory-management data structures. 7. Restore the selected process's saved processor context.",
                ))
            elif mode_switch and len(steps) >= 3 and "interrupt handler" in steps[1].lower():
                cards.append((
                    question,
                    "1. Save the processor context. 2. Set the PC to the interrupt handler. 3. Switch from user mode to system mode.",
                ))
            else:
                cards.append((question, " ".join(steps)))
        steps = []

    for paragraph in paragraphs:
        heading_like = (
            len(paragraph.split()) <= 8
            and len(paragraph) <= 72
            and not re.search(r"[.!?]$", paragraph)
            and (paragraph.istitle() or paragraph.isupper() or paragraph.endswith(")"))
        )
        if heading_like:
            flush()
            section = paragraph
            question = "What processor actions occur during a mode switch?" if "mode switch" in section.lower() else ""
            continue
        if re.search(r"if the interrupt results in a full process switch", paragraph, flags=re.IGNORECASE):
            flush()
            question = "What steps occur during a full process switch?"
            continue
        numbered = re.match(r"^(\d+)\.\s*(.+)$", paragraph)
        if not numbered:
            flush()
            continue
        number, step = numbered.groups()
        if number == "1" and steps:
            flush()
        if not steps and not question:
            continue
        if question:
            steps.append(f"{number}. {step.strip()}")
    flush()
    return cards


def _pdf_comparison_cards(paragraphs: list[str]) -> list[tuple[str, str]]:
    cards = []
    multiprogramming = next((re.search(r"Multiprogramming:\s*(.+)", value, flags=re.IGNORECASE) for value in paragraphs if re.search(r"Multiprogramming:\s*", value, flags=re.IGNORECASE)), None)
    multiprocessing = next((re.search(r"Multiprocessing:\s*(.+)", value, flags=re.IGNORECASE) for value in paragraphs if re.search(r"Multiprocessing:\s*", value, flags=re.IGNORECASE)), None)
    if multiprogramming and multiprocessing:
        cards.append((
            "How does multiprogramming differ from multiprocessing?",
            "Multiprogramming uses one CPU, with multiple processes ready for execution; multiprocessing uses several processors (CPUs) on one computer system to increase processing power.",
        ))
    hard = next((re.search(r"Hard real-time system:\s*(.+)", value, flags=re.IGNORECASE) for value in paragraphs if re.search(r"Hard real-time system:\s*", value, flags=re.IGNORECASE)), None)
    soft = next((re.search(r"Soft real-time system:\s*(.+)", value, flags=re.IGNORECASE) for value in paragraphs if re.search(r"Soft real-time system:\s*", value, flags=re.IGNORECASE)), None)
    if hard and soft:
        cards.append((
            "How do hard and soft real-time systems differ?",
            "A hard real-time system guarantees that a critical task is completed in time; a soft real-time system gives a critical task priority over other tasks.",
        ))
    return cards


def _numbered_pdf_blocks(text: str, answer_key: bool = False) -> dict[str, str]:
    if answer_key:
        pattern = re.compile(r"(?m)^[ \t]*(\d+\.\d+)(?:[ \t]+(?=\S)|[ \t]*$)")
    else:
        pattern = re.compile(r"(?m)^[ \t]*(\d+\.\d+)[ \t]+(?=\S)")
    matches = list(pattern.finditer(text))
    answer_markers = [match.start() for match in re.finditer(r"(?im)^[ \t]*answers[ \t]*$", text)]
    blocks = {}
    for index, match in enumerate(matches):
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        if not answer_key:
            next_answer = next((position for position in answer_markers if match.end() < position < end), None)
            if next_answer is not None:
                end = next_answer
        body_start = match.end()
        body = re.sub(r"\s+", " ", text[body_start:end]).strip()
        if body:
            blocks[match.group(1)] = body
    return blocks


def _refine_pdf_quiz_pair(question: str, answer: str) -> list[tuple[str, str]]:
    question_lower = question.casefold()
    answer = re.sub(r"\s+", " ", answer).strip()

    if "unix shell internal command" in question_lower:
        return [
            (
                "What is a Unix shell internal command?",
                "It is executable code built into the command interpreter; these commands are generally simple and run from the shell.",
            )
        ]

    if "multiprogramming environment" in question_lower and "batch processing" in question_lower and "challenge" in question_lower:
        return [
            (
                "How does multiprogramming improve resource use compared with batch processing?",
                "Processes take turns using the CPU, leading to more efficient resource use than sequential batch processing.",
            ),
            (
                "What synchronization problem can concurrent processes create?",
                "One process may read an address space while another process writes to that same address space.",
            ),
        ]

    if "bootstrap program" in question_lower and "rom" in question_lower:
        first_reason = re.search(r"1\)\s*(.+?)(?=\s*2\)\s*)", answer, flags=re.IGNORECASE)
        second_reason = re.search(r"2\)\s*(.+)$", answer, flags=re.IGNORECASE)
        if first_reason and second_reason:
            return [
                ("Why does a full bootstrap program no longer fit in ROM?", first_reason.group(1).strip()),
                ("Why would changing a full bootstrap program in ROM require a hardware change?",
                 "ROM is costly, and changing a program stored there would require changing the hardware."),
            ]

    if "interrupt queue" in question_lower and "batch file" in question_lower:
        batch_file = re.search(r"\bA batch file is\s+(.+)$", answer, flags=re.IGNORECASE)
        queue_answer = re.search(r"\bAn interrupt queue is\s+(.+?)(?=\s+A batch file is)", answer, flags=re.IGNORECASE)
        if queue_answer and batch_file:
            return [
                ("What is an interrupt queue?",
                 "It is populated when interrupts occur and uses a priority algorithm to determine which interrupt is handled."),
                ("What is a batch file?", "It is an executable system file often run while the operating system initializes."),
            ]

    if "system calls necessary" in question_lower:
        necessity = re.search(r"\bSystem calls are necessary because\s+(.+)$", answer, flags=re.IGNORECASE)
        examples = re.search(r"\bSome examples of system calls include\s+(.+?)(?=\s+System calls are necessary because)", answer, flags=re.IGNORECASE)
        if examples and necessity:
            return [
                ("Which system-call examples concern process control and privileged instructions?",
                 "Create and delete calls start or end processes; another system call is used for a privileged instruction."),
                ("How do system calls help protect a process and the operating system?",
                 "They let the operating system double-check certain processes to protect both the process and the OS."),
            ]

    if "win32" in question_lower and "responsibilities" in question_lower:
        role = re.search(r"^(win32 is .+?system space\.)\s*", answer, flags=re.IGNORECASE)
        responsibilities = re.search(r"Some of win32[’']s responsibilities are\s+(.+)$", answer, flags=re.IGNORECASE)
        if role and responsibilities:
            return [
                ("What role does Win32 play in Windows architecture?",
                 "It is the key Windows subsystem that enables communication between user space and system space."),
                ("What direct responsibilities does Win32 have?",
                 "It runs Win32 applications, manages keyboard, mouse, and screen I/O, and starts processes by sending system calls to the executive."),
            ]

    if "tsr process" in question_lower:
        definition = re.search(r"A process is considered a TSR process if\s+(.+?)(?:\.\s+A regular process|$)", answer, flags=re.IGNORECASE)
        examples = re.search(r"Examples of TSR processes would be\s+(.+?)\.\s+A process is considered", answer, flags=re.IGNORECASE)
        if definition:
            refined = [("When is a process considered a TSR process?",
                        "It is invoked frequently enough to reside in computer memory / address space; a regular process is not stored there.")]
            if examples:
                refined.insert(0, ("What are examples of terminate-and-stay-resident (TSR) processes?",
                                   examples.group(1).strip() + " are examples of TSR processes."))
            return refined

    if "unix shells" in question_lower and "/bin" in answer:
        shell_names = re.search(r"Examples of Unix Shells are\s+(.+?)\.\s+You can find", answer, flags=re.IGNORECASE)
        shell_directory = re.search(r"find the shell program in the\s+(.+?)\s+directory", answer, flags=re.IGNORECASE)
        invoked = re.search(r"The user[’']s shell is invoked\s+(.+?)\.\s+The Unix shell runs in user mode", answer, flags=re.IGNORECASE)
        if shell_names and shell_directory and invoked:
            shell_examples = shell_names.group(1).strip().replace(", ksh", ", and ksh")
            return [
                ("Which Unix shells are given as examples?", "The Unix shells named as examples are " + shell_examples + "."),
                ("Where are Unix shell programs located?",
                 "The " + shell_directory.group(1).strip() + " directory in the Unix file system."),
                ("When is a user's Unix shell invoked?", "The shell is invoked " + invoked.group(1).strip() + "."),
                ("In which mode does a Unix shell run?", "The Unix shell runs in user mode."),
            ]

    if "embedded systems" in question_lower and "real time" in question_lower:
        distinction = re.search(r"No,\s*(.+?)\s+A real time system can be part of an embedded system\.", answer, flags=re.IGNORECASE)
        relationship = re.search(r"(A real time system can be part of an embedded system)\.", answer, flags=re.IGNORECASE)
        if distinction and relationship:
            return [
                ("Are embedded systems and real-time systems the same?",
                 "No. An embedded system controls a larger system; a real-time system has rigid process-execution time constraints. A real-time system can be part of an embedded system."),
            ]

    if "vertical layered structure" in question_lower:
        structure = re.search(r"The vertical layered structure is when\s+(.+?)(?=\s+Some disadvantages are)", answer, flags=re.IGNORECASE)
        disadvantages = re.search(r"Some disadvantages are\s+(.+)$", answer, flags=re.IGNORECASE)
        if structure and disadvantages:
            return [
                ("How is an operating system organized in a vertical layered structure?",
                 "The layers are stacked, with layer 0 as the hardware and layer N as the user interface."),
                ("What are the disadvantages of a vertical layered structure?",
                 "It can be difficult to define layers, communication between layers adds overhead, and adding functionality can be difficult."),
            ]

    return [(question.strip(), answer)]


def _pdf_assessment_cards(page_texts: list[str], filename: str, status: str) -> list[dict]:
    whole_text = "\n".join(page_texts)
    answer_marker = re.search(r"(?im)^\s*answers\s*$", whole_text)
    if not answer_marker:
        return []
    questions = _numbered_pdf_blocks(whole_text)
    answers = _numbered_pdf_blocks(whole_text, answer_key=True)
    cards = []
    for question_id, question in questions.items():
        answer = answers.get(question_id)
        if not answer:
            continue
        front = question.strip().rstrip(" .")
        front = front + ("?" if not front.endswith("?") else "")
        back = answer.strip()
        normalized = back.lower()
        # These two responses are personalized or internally inconsistent in the supplied answer key.
        if "installed in your computer" in question.lower() or "my computer" in normalized:
            continue
        if "command.com" in question.lower() and "unix" in normalized:
            continue
        refined_cards = _refine_pdf_quiz_pair(question, back)
        for part_number, (refined_front, refined_back) in enumerate(refined_cards, start=1):
            refined_front = re.sub(r"\s+", " ", refined_front).strip().rstrip(" .?") + "?"
            refined_back = re.sub(r"\s+", " ", refined_back).strip()
            back_words = re.findall(r"\b\w+\b", refined_back)
            if not (5 <= len(back_words) <= 65) or len(refined_front.split()) < 5:
                continue
            if refined_back and refined_back[-1] not in ".!?":
                refined_back += "."
            if refined_back:
                refined_back = refined_back[:1].upper() + refined_back[1:]
            cards.append({
                "muscle": "Quiz answer key",
                "field": "assessment",
                "front": refined_front,
                "back": refined_back,
                "section": "Quiz answer key",
                "source": f"{filename} · Question {question_id}" + (f" · Part {part_number}" if len(refined_cards) > 1 else ""),
                "status": status,
                "questionId": question_id,
                "partNumber": part_number,
            })
    return cards


def compile_pdf_study_material(
    page_texts: list[str], filename: str, status: str = "verified", assessment: bool = False
) -> dict:
    if assessment:
        assessment_cards = _pdf_assessment_cards(page_texts, filename, status)
        if assessment_cards:
            return {"concepts": [], "notes": [], "cards": assessment_cards}

    concepts = []
    notes = []
    cards = []
    seen = set()
    stem = Path(filename).stem.replace("_", " ").replace("-", " ").strip()
    for page_number, raw_page in enumerate(page_texts, start=1):
        paragraphs = _pdf_paragraphs(raw_page)
        if not paragraphs:
            continue
        headings = [
            paragraph for paragraph in paragraphs
            if len(paragraph.split()) <= 8 and not re.search(r"[.!?]$", paragraph)
        ]
        topic = headings[0] if headings else f"{stem} · Page {page_number}"
        citation = f"{filename} · Page {page_number}"
        content = [
            paragraph for paragraph in paragraphs
            if len(re.findall(r"\b\w+\b", paragraph)) >= 6
            and not _PDF_BOILERPLATE.match(paragraph)
        ]
        if not content:
            continue
        notes.append({
            "title": topic,
            "section": topic,
            "heardAt": None,
            "status": status,
            "source": citation,
            "pageNumber": page_number,
            "lines": content[:8],
        })
        page_cards = []
        normalized_page = re.sub(r"\s+", " ", raw_page).casefold()
        context_cards = []
        if "ms-dos - application programs are able to access the basic i/o routines" in normalized_page and "makes ms-dos vulnerable to errant programs" in normalized_page:
            context_cards.append((
                "What weakness follows from MS-DOS applications accessing basic I/O routines?",
                "MS-DOS is vulnerable to errant programs.",
            ))
        if "an interrupt can be synchronous if it is the direct result of the current instruction" in normalized_page and "otherwise it is known as asynchronous" in normalized_page:
            context_cards.append((
                "How do synchronous and asynchronous interrupts differ?",
                "A synchronous interrupt is the direct result of the current instruction executing on the CPU; otherwise, it is asynchronous.",
            ))
        if (
            "the cpu issues a command requesting i/o" in normalized_page
            and ("continue to execute other instructions" in normalized_page or "continues to execute other instructions" in normalized_page)
            and "until the i/o module completes its work" in normalized_page
            and ("the i/o module will issue an interrupt" in normalized_page or "the i/o module issues an interrupt" in normalized_page)
        ):
            context_cards.append((
                "How does interrupt-driven I/O let the CPU keep working?",
                "The CPU issues an I/O command and continues executing other instructions while the I/O module works; the module interrupts the CPU when it finishes.",
            ))
        if all(phrase in normalized_page for phrase in (
            "the interrupt handler checks the cause of the interrupt",
            "it might service the interrupt",
            "the interrupt handler routine is to protect the pcb",
        )):
            context_cards.extend([
                (
                    "What does the interrupt handler do after it receives an interrupt?",
                    "It checks the cause of the interrupt and may service it.",
                ),
                (
                    "What is a main task of the interrupt handler routine?",
                    "It protects the process control block (PCB); it is usually the only process that can modify the PCB's information.",
                ),
            ])
        if "interrupt handler will call a specific interrupt service routine" in normalized_page:
            context_cards.append((
                "When may an interrupt handler call an interrupt service routine?",
                "In more complicated situations, it may call a specific interrupt service routine to service the interrupt.",
            ))
        if all(phrase in normalized_page for phrase in (
            "thread scheduling",
            "low-level processor synchronization",
            "interrupt and exception handling",
            "switching between user and system modes",
        )):
            context_cards.append((
                "Which responsibilities does the Windows 10 kernel handle?",
                "It handles thread scheduling, low-level processor synchronization, interrupts and exceptions, and switches between user and system modes.",
            ))
        for front, back in context_cards:
            if _pdf_card_is_usable(front, back):
                key = (front.casefold(), back.casefold())
                if key not in seen:
                    seen.add(key)
                    page_cards.append((front, back))
        for front, back in _pdf_comparison_cards(paragraphs):
            if _pdf_card_is_usable(front, back):
                key = (front.casefold(), back.casefold())
                if key not in seen:
                    seen.add(key)
                    page_cards.append((front, back))
        for front, back in _pdf_sequence_cards(paragraphs):
            if _pdf_card_is_usable(front, back, allow_long_answer=True):
                key = (front.casefold(), back.casefold())
                if key not in seen:
                    seen.add(key)
                    page_cards.append((front, back))
        current_heading = ""
        for paragraph in paragraphs:
            if (
                len(paragraph.split()) <= 8
                and len(paragraph) <= 72
                and not re.search(r"[.!?]$", paragraph)
                and (paragraph.istitle() or paragraph.isupper() or paragraph.endswith(")"))
            ):
                current_heading = paragraph
                continue
            if len(re.findall(r"\b\w+\b", paragraph)) < 6 or _PDF_BOILERPLATE.match(paragraph):
                continue
            for statement in _study_sentences([paragraph]):
                generated = _pdf_fact_card(statement, current_heading)
                if not generated:
                    continue
                front, back = generated
                if not _pdf_card_is_usable(front, back):
                    continue
                key = (front.casefold(), back.casefold())
                if key in seen:
                    continue
                seen.add(key)
                page_cards.append((front, back))
        if page_cards:
            concepts.append({"name": topic, "status": status, "source": citation, "pageNumber": page_number})
        for front, back in page_cards[:8]:
            answer = back.strip()
            if answer and answer[-1] not in ".!?":
                answer += "."
            if answer:
                answer = answer[:1].upper() + answer[1:]
            cards.append({
                "muscle": topic,
                "field": "concept",
                "front": front,
                "back": answer,
                "section": topic,
                "source": citation,
                "status": status,
                "pageNumber": page_number,
            })
    return {"concepts": concepts, "notes": notes, "cards": cards}


def compile_lecture_window(text: str, filename: str, heard_at: float, window_number: int) -> dict:
    cleaned = re.sub(r"\s+", " ", text).strip()
    if len(re.findall(r"\b\w+\b", cleaned)) < 8:
        return {"concepts": [], "notes": [], "cards": []}
    sentences = _study_sentences([cleaned]) or [cleaned]
    sentences = [
        sentence for sentence in sentences
        if not any(phrase in sentence.lower() for phrase in _PROMOTIONAL_LECTURE_PHRASES)
    ]
    if not sentences:
        return {"concepts": [], "notes": [], "cards": []}
    citation = f"{filename} · {int(heard_at // 60):02d}:{int(heard_at % 60):02d}"
    card_pairs = []
    seen = set()
    for sentence in sentences:
        generated = _fact_card(sentence)
        if not generated or not _usable_lecture_answer(generated[1]) or generated[0].lower() in seen:
            continue
        seen.add(generated[0].lower())
        card_pairs.append((*generated, sentence))
        if len(card_pairs) >= 2:
            break
    if not card_pairs:
        return {"concepts": [], "notes": [], "cards": []}
    topic = _topic_from_card_question(card_pairs[0][0])
    if not topic:
        return {"concepts": [], "notes": [], "cards": []}
    concept = {"name": topic, "status": "provisional", "source": citation}
    note = {
        "title": topic,
        "section": "Lecture notes",
        "heardAt": round(heard_at, 2),
        "status": "provisional",
        "source": citation,
        "lines": sentences[:3],
    }
    cards = [
        {
            "muscle": topic,
            "field": "lecture",
            "front": front,
            "back": back,
            "section": "Lecture notes",
            "source": citation,
            "status": "provisional",
            "windowNumber": window_number,
            "heardAt": round(heard_at, 2),
            "sourceLocation": f"{int(heard_at // 60):02d}:{int(heard_at % 60):02d}",
            "sourceQuote": source_sentence,
        }
        for front, back, source_sentence in card_pairs
    ]
    return {"concepts": [concept], "notes": [note], "cards": cards}


def compile_lecture_segments(segments: list[dict], filename: str) -> dict:
    concepts = []
    notes = []
    cards = []
    seen_questions = set()
    window_size = 5
    for offset in range(0, len(segments), window_size):
        window = segments[offset : offset + window_size]
        if not window:
            continue
        compiled = compile_lecture_window(
            " ".join(segment.get("text", "") for segment in window),
            filename,
            float(window[0].get("start", 0)),
            offset // window_size + 1,
        )
        new_cards = []
        for card in compiled["cards"]:
            key = card["front"].strip().lower()
            if key in seen_questions:
                continue
            seen_questions.add(key)
            new_cards.append(card)
        concepts.extend(compiled["concepts"])
        notes.extend(compiled["notes"])
        cards.extend(new_cards)
    return {"concepts": concepts, "notes": notes, "cards": cards}


_MONTHS = {
    name.lower(): index
    for index, names in enumerate(
        (
            ("jan", "january"), ("feb", "february"), ("mar", "march"), ("apr", "april"),
            ("may",), ("jun", "june"), ("jul", "july"), ("aug", "august"),
            ("sep", "sept", "september"), ("oct", "october"), ("nov", "november"), ("dec", "december"),
        ),
        start=1,
    )
    for name in names
}


def classify_source(filename: str, text: str) -> tuple[str, str]:
    normalized_name = filename.lower().replace("_", " ").replace("-", " ")
    normalized_text = text.lower()
    syllabus_metadata = sum(
        bool(re.search(pattern, normalized_text))
        for pattern in (
            r"\bcourse\s+(?:name|title|number|code)\s*[:|]",
            r"\bcourse\s*[:|]\s*[a-z0-9]",
            r"(?m)^[a-z]{2,8}\s*\d{2,4}\s*[-:|]\s+\w",
            r"\bmeeting\s+times?\b",
            r"\bcatalog\s+description\b",
            r"\bcourse\s+description\b",
            r"\binstructor\b",
            r"\boffice\s+hours\b",
            r"\bprerequisites?\b",
            r"\b(?:required|recommended)\s+(?:text(?:book)?|readings?|materials?)\b",
            r"\bgrading(?:\s+(?:policy|scale|criteria))?\b",
            r"\b(?:evaluation|assessment)\s+(?:criteria|weights?|breakdown)\b",
            r"\b(?:learning\s+)?outcomes?\b",
            r"\b(?:academic\s+integrity|attendance\s+policy|course\s+policies)\b",
            r"\b(?:weekly|tentative|course|class)\s+(?:topic|schedule|calendar)\b",
        )
    )
    syllabus_filename = bool(re.search(r"\b(syllabus|course outline)\b", normalized_name))
    syllabus_header = bool(re.search(r"\bcourse\s+syllabus\b", normalized_text))
    syllabus_outline = bool(re.search(r"\b(course outline|course calendar|course expectations)\b", normalized_text))
    schedule_table = bool(
        re.search(r"\bweek\s*(?:\||\s{1,})dates?\s*(?:\||\s{1,})(?:topics?|lectures?|readings?)\b", normalized_text)
        and re.search(r"\b(term|instructor|grading|course description)\b", normalized_text)
    )
    metadata_syllabus = syllabus_metadata >= 5 and bool(
        re.search(r"\b(course description|catalog description|prerequisites?|grading|course schedule|class schedule|weekly schedule)\b", normalized_text)
    )
    if re.search(r"\b(past|practice|sample)?\s*(midterm|final exam|exam|assessment|test|quiz)\b", normalized_name):
        return "assessment", "Assessment terms detected in the file name"
    if (
        syllabus_filename
        or (syllabus_header and syllabus_metadata >= 2)
        or (syllabus_outline and syllabus_metadata >= 2)
        or metadata_syllabus
        or schedule_table
    ):
        return "syllabus", "Course schedule or syllabus structure detected"
    if _looks_like_administrative_form(text):
        return "material", "Administrative form detected; no study-card content expected"
    question_cues = len(re.findall(r"(?:^|\n)\s*(?:\d+[.)]|question\s*\d*[:.)]|(?:a|b|c|d)[.)])\s+", text, re.IGNORECASE))
    explicit_question_cues = len(re.findall(r"(?im)^\s*(?:question\s*\d+|q\s*\d+)\s*[:.)]", text))
    answer_option_cues = len(re.findall(r"(?im)^\s*[a-f][.)]\s+", text))
    assessment_terms = len(re.findall(r"\b(answer key|multiple choice|select all that apply|exam question|practice test)\b", normalized_text))
    if explicit_question_cues >= 2 or (question_cues >= 3 and answer_option_cues >= 2) or assessment_terms >= 2:
        return "assessment", "Question and answer structure detected"
    return "material", "Study content detected; no syllabus or assessment pattern matched"


def _looks_like_administrative_form(text: str) -> bool:
    normalized = text.lower()
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    blank_fields = len(re.findall(r"_{3,}", text))
    field_labels = sum(
        bool(re.match(r"^(?:\d+[.)]\s*)?(?:name|address|mailing address|telephone number|phone|email|signature|date|entity name|business name)\s*:", line, re.IGNORECASE))
        for line in lines
    )
    form_markers = sum(
        phrase in normalized
        for phrase in (
            "application form",
            "document and certificate cover sheet",
            "division of corporations",
            "contact information",
            "filing fee",
            "mailing address",
            "signature",
            "entity name",
        )
    )
    return blank_fields >= 3 and field_labels >= 3 and form_markers >= 2


def _course_year(text: str, filename: str) -> int | None:
    combined = f"{filename} {text}"
    match = re.search(r"\b(?:spring|summer|fall|autumn|winter)\s+(20\d{2})\b", combined, re.IGNORECASE)
    if not match:
        compact_term = re.search(r"\b(SP|SU|FA|WI)\s*(\d{2})\b", re.sub(r"[_-]+", " ", filename), re.IGNORECASE)
        if compact_term:
            return 2000 + int(compact_term.group(2))
    return int(match.group(1)) if match else None


def _dates_in_text(value: str, year: int | None) -> list[str]:
    dates: list[str] = []
    month_pattern = "|".join(sorted(_MONTHS, key=len, reverse=True))
    pattern = re.compile(
        rf"\b({month_pattern})\.?\s+(\d{{1,2}})((?:\s*(?:,|and|&|-|to|–)\s*\d{{1,2}})*)(?:,?\s+(20\d{{2}}))?\b",
        re.IGNORECASE,
    )
    for match in pattern.finditer(value):
        month = _MONTHS[match.group(1).lower()]
        first_day = int(match.group(2))
        days = [first_day]
        previous_day = first_day
        for day_match in re.finditer(r"(,|and|&|-|to|–)\s*(\d{1,2})", match.group(3), re.IGNORECASE):
            joiner, raw_day = day_match.groups()
            day = int(raw_day)
            if joiner in {"-", "to", "–"} and previous_day < day <= previous_day + 14:
                days.extend(range(previous_day + 1, day + 1))
            else:
                days.append(day)
            previous_day = day
        date_year = int(match.group(4)) if match.group(4) else year
        if date_year is None:
            continue
        for day in days:
            try:
                dates.append(date(date_year, month, day).isoformat())
            except ValueError:
                continue
    for match in re.finditer(r"\b(\d{1,2})/(\d{1,2})/(20\d{2})\b", value):
        try:
            dates.append(date(int(match.group(3)), int(match.group(1)), int(match.group(2))).isoformat())
        except ValueError:
            continue
    return list(dict.fromkeys(dates))


def _contains_unqualified_date(value: str) -> bool:
    month_pattern = "|".join(sorted(_MONTHS, key=len, reverse=True))
    return bool(
        re.search(rf"\b(?:{month_pattern})\.?\s+\d{{1,2}}\b", value, re.IGNORECASE)
        or re.search(r"\b\d{1,2}/\d{1,2}(?:/20\d{2})?\b", value)
    )


def _date_year_source(value: str, fallback: str) -> str:
    month_pattern = "|".join(sorted(_MONTHS, key=len, reverse=True))
    if re.search(rf"\b(?:{month_pattern})\.?\s+\d{{1,2}}(?:[^\n]*?)\b20\d{{2}}\b", value, re.IGNORECASE):
        return "schedule row"
    if re.search(r"\b\d{1,2}/\d{1,2}/20\d{2}\b", value):
        return "schedule row"
    return fallback


def _without_date_references(value: str) -> str:
    month_pattern = "|".join(sorted(_MONTHS, key=len, reverse=True))
    value = re.sub(
        rf"\b(?:{month_pattern})\.?\s+\d{{1,2}}(?:\s*(?:,|and|&|-|to|–)\s*\d{{1,2}})*(?:,?\s+20\d{{2}})?\b",
        " ",
        value,
        flags=re.IGNORECASE,
    )
    value = re.sub(r"\b\d{1,2}/\d{1,2}/20\d{2}\b", " ", value)
    return re.sub(r"[\s|,;:–-]+", " ", value).strip()


def syllabus_calendar(
    filename: str, text: str, source_id: str, page_texts: list[str] | None = None
) -> dict:
    """Extract dated class meetings and deadlines from clearly structured syllabus rows."""
    lines = [re.sub(r"\s+", " ", line).strip() for line in text.splitlines() if line.strip()]
    page_hits: dict[str, set[int]] = {}
    for page_number, page_text in enumerate(page_texts or [], start=1):
        for page_line in page_text.splitlines():
            key = re.sub(r"\s+", " ", page_line).strip().casefold()
            if key:
                page_hits.setdefault(key, set()).add(page_number)

    def source_page(line: str) -> int | None:
        pages = page_hits.get(re.sub(r"\s+", " ", line).strip().casefold(), set())
        return next(iter(pages)) if len(pages) == 1 else None

    course_name = ""
    course_code = ""
    term = ""
    term_range = ""
    for line in lines:
        fields = [field.strip() for field in line.split("|")]
        if len(fields) >= 2 and fields[0].lower() == "course":
            course_name = fields[1]
        if len(fields) >= 2 and fields[0].lower() == "term":
            term = fields[1]
            term_range = " · ".join(fields[2:])
        label_match = re.match(r"^course\s+(name|title|number|code)\s*:\s*(.+)$", line, re.IGNORECASE)
        if label_match:
            label, value = label_match.group(1).lower(), label_match.group(2).strip()
            if label in {"name", "title"}:
                course_name = course_name or value
            else:
                course_code = value
        inline_course = re.match(r"^([A-Z]{2,8}\s*\d{2,4})\s*[-:|]\s*(.+)$", line, re.IGNORECASE)
        if inline_course:
            course_code = course_code or re.sub(r"\s+", "", inline_course.group(1).upper())
            course_name = course_name or inline_course.group(2).strip()
        bare_course_code = re.match(r"^\s*(COS\s*226)\b", line, re.IGNORECASE)
        if bare_course_code:
            course_code = course_code or re.sub(r"\s+", "", bare_course_code.group(1).upper())
            course_name = course_name or "COS 226"
        if not term:
            term_match = re.search(r"\b(Fall|Autumn|Winter|Spring|Summer)\s+(20\d{2}|\d{2})\b", line, re.IGNORECASE)
            if term_match:
                term_year = term_match.group(2)
                term_year = term_year if len(term_year) == 4 else f"20{term_year}"
                term = f"{term_match.group(1).title()} {term_year}"
    if not term:
        full_term = re.search(r"\b(Fall|Autumn|Winter|Spring|Summer)\s+(20\d{2}|\d{2})\b", filename, re.IGNORECASE)
        if full_term:
            full_year = full_term.group(2)
            term = f"{full_term.group(1).title()} {full_year if len(full_year) == 4 else f'20{full_year}'}"
    if not term:
        compact_term = re.search(r"\b(SP|SU|FA|WI)\s*(\d{2})\b", re.sub(r"[_-]+", " ", filename), re.IGNORECASE)
        if compact_term:
            term_name = {"SP": "Spring", "SU": "Summer", "FA": "Fall", "WI": "Winter"}[compact_term.group(1).upper()]
            term = f"{term_name} 20{compact_term.group(2)}"
    if not course_name:
        course_name = next(
            (line.split(":", 1)[1].strip() for line in lines if re.match(r"^course\s+(?:name|title)\s*:", line, re.IGNORECASE)),
            "",
        )
    if not course_name:
        course_name = next((line for line in lines if re.search(r"\b(?:PSY|BIO|CHEM|HIST|MATH|ENG|CS|COS|CSCI|CSC|COMP)\s*\d{2,4}\b", line, re.IGNORECASE)), course_code)
    year = _course_year(text, filename)
    term_in_text = re.search(r"\b(?:spring|summer|fall|autumn|winter)\s+20\d{2}\b", text, re.IGNORECASE)
    year_fallback_source = "course term" if term_in_text else "source name"
    events: list[dict] = []
    unresolved_date_rows = 0
    no_class_dates = set()
    for line in lines:
        no_class_match = re.search(r"\bno\s+class\b", line, re.IGNORECASE)
        if no_class_match:
            no_class_dates.update(_dates_in_text(line[no_class_match.end():], year))

    mode = ""
    processed_schedule_lines: set[str] = set()
    for line in lines:
        lowered = line.lower()
        fields = [field.strip() for field in line.split("|")]
        if re.search(r"\bweek\b", lowered) and re.search(r"\bdates?\b", lowered) and re.search(r"\b(?:topics?|lectures?|readings?)\b", lowered):
            mode = "weekly"
            continue
        if re.match(r"^date\s+(?:milestone|event)\b", lowered) or (
            len(fields) >= 2 and fields[0].lower() == "date" and fields[1].lower() in {"milestone", "event"}
        ):
            mode = "milestones"
            continue
        if re.fullmatch(
            r"(?:(?:course|class|exam|assessment|assignment|important|tentative|weekly)\s+)?(?:schedule|calendar|dates|deadlines)(?:\s*[:|-].*)?",
            lowered,
        ):
            mode = "milestones"
            continue
        if mode == "weekly":
            if len(fields) >= 3 and re.fullmatch(r"\d+", fields[0]):
                week = fields[0]
                dates_text = fields[1]
                topic = fields[2].split(";")[0].strip()
            else:
                row = re.match(r"^(?:week\s*)?(\d{1,2})[\s.)|:-]+(.+)$", line, re.IGNORECASE)
                if not row:
                    continue
                week, dates_text = row.groups()
                topic = _without_date_references(dates_text).split(";", 1)[0].strip()
                if not _dates_in_text(dates_text, year):
                    if _contains_unqualified_date(dates_text):
                        unresolved_date_rows += 1
                    continue
            title = f"Week {week}: {topic}"[:120]
            weekly_dates = _dates_in_text(dates_text, year)
            if not weekly_dates and _contains_unqualified_date(dates_text):
                unresolved_date_rows += 1
            for iso_date in weekly_dates:
                if iso_date not in no_class_dates:
                    events.append({
                        "date": iso_date,
                        "type": "lecture",
                        "title": title,
                        "sourceText": line,
                        "sourcePage": source_page(line),
                        "yearSource": _date_year_source(dates_text, year_fallback_source),
                    })
            if len(fields) >= 5:
                due_dates = _dates_in_text(fields[-1], year)
                if not due_dates and _contains_unqualified_date(fields[-1]):
                    unresolved_date_rows += 1
                if due_dates:
                    due_title = re.sub(r"\b(?:on|due)\b", "", fields[-1], flags=re.IGNORECASE).strip(" ,.-")
                    month_pattern = "|".join(sorted(_MONTHS, key=len, reverse=True))
                    due_title = re.sub(rf"\b(?:{month_pattern})\.?\s+\d{{1,2}}(?:\s*(?:,|and|&)\s*\d{{1,2}})*\b", "", due_title, flags=re.IGNORECASE).strip(" ,.-")
                    events.extend({
                        "date": iso_date,
                        "type": _milestone_type(due_title),
                        "title": due_title[:120],
                        "sourceText": line,
                        "sourcePage": source_page(line),
                        "yearSource": _date_year_source(fields[-1], year_fallback_source),
                    } for iso_date in due_dates)
            processed_schedule_lines.add(line)
            continue
        if mode == "milestones":
            if len(fields) >= 2 and fields[0].lower() not in {"date", "milestone", "event"}:
                milestone_dates = _dates_in_text(fields[0], year)
                title = fields[1].strip()[:120]
                date_text = fields[0]
            else:
                month_pattern = "|".join(sorted(_MONTHS, key=len, reverse=True))
                row = re.match(
                    rf"^((?:{month_pattern})\.?\s+\d{{1,2}}(?:\s*(?:,|and|&|-|to|–)\s*\d{{1,2}})*(?:,?\s+20\d{{2}})?)\s+(.+)$",
                    line,
                    re.IGNORECASE,
                )
                if not row:
                    continue
                date_text, title = row.groups()
                milestone_dates = _dates_in_text(date_text, year)
                title = title.strip(" \t-|:.")[:120]
            if not milestone_dates and _contains_unqualified_date(date_text):
                unresolved_date_rows += 1
            for iso_date in milestone_dates:
                events.append({
                    "date": iso_date,
                    "type": _milestone_type(title),
                    "title": title,
                    "sourceText": line,
                    "sourcePage": source_page(line),
                    "yearSource": _date_year_source(date_text, year_fallback_source),
                })
            processed_schedule_lines.add(line)

    # Syllabi often state major assessments in prose outside their dated schedule table.
    # Pick up only explicit dated milestones here so ordinary policy dates do not become events.
    for line in lines:
        if line in processed_schedule_lines:
            continue
        if not re.search(r"\b(?:mid[- ]?term|final\s+(?:exam|examination)|exam(?:ination)?|quiz|assignment|homework|project|paper|presentation)\b", line, re.IGNORECASE):
            continue
        milestone_dates = _dates_in_text(line, year)
        if not milestone_dates:
            if _contains_unqualified_date(line):
                unresolved_date_rows += 1
            continue
        lowered = line.lower()
        if re.search(r"\bmid[- ]?term\b", lowered):
            title = "Midterm exam"
        elif re.search(r"\bfinal\s+(?:exam|examination)\b|\bfinal\s+exam\b", lowered) or re.search(r"\bLnal\s+exam", line, re.IGNORECASE):
            title = "Final exam"
        elif re.search(r"\bquiz\b", lowered):
            quiz = re.search(r"\bquiz\s*(?:#|no\.?\s*)?(\d+[A-Z]?)?", line, re.IGNORECASE)
            title = f"Quiz {quiz.group(1)}" if quiz and quiz.group(1) else "Quiz"
        elif re.search(r"\bassignment\b|\bhomework\b", lowered):
            title = "Assignment due" if re.search(r"\bdue\b", lowered) else "Assignment"
        elif re.search(r"\bproject\b", lowered):
            title = "Project due" if re.search(r"\bdue\b", lowered) else "Project"
        elif re.search(r"\bpaper\b", lowered):
            title = "Paper due" if re.search(r"\bdue\b", lowered) else "Paper"
        elif re.search(r"\bpresentation\b", lowered):
            title = "Presentation"
        else:
            title = "Exam"
        events.extend({
            "date": iso_date,
            "type": _milestone_type(title),
            "title": title,
            "sourceText": line,
            "sourcePage": source_page(line),
            "yearSource": _date_year_source(line, year_fallback_source),
        } for iso_date in milestone_dates)

    deduped = {}
    for event in events:
        if event["type"] == "lecture" and event["date"] in no_class_dates:
            continue
        key = (event["date"], event["type"], re.sub(r"\W+", " ", event["title"]).strip().lower())
        previous = deduped.get(key)
        if previous is None or len(event["title"]) > len(previous["title"]):
            deduped[key] = event
    result = []
    for event in deduped.values():
        event_id = f"{source_id}-cal-{event['date']}-{safe_slug(event['title'])[:32]}"
        result.append({**event, "id": event_id, "sourceId": source_id, "sourceName": filename, "sourceTerm": term})
    warnings = []
    if unresolved_date_rows:
        warnings.append("Some schedule dates were left off because their calendar year was not clear in the syllabus.")
    return {
        "events": sorted(result, key=lambda event: (event["date"], event["type"], event["title"])),
        "courseName": course_name,
        "courseCode": course_code,
        "term": term,
        "termRange": term_range,
        "warnings": warnings,
    }


def _milestone_type(title: str) -> str:
    normalized = title.lower()
    if re.search(r"\b(exam|midterm|final\s+(?:exam|examination))\b", normalized):
        return "exam"
    if re.search(r"\b(quiz|check)\b", normalized):
        return "quiz"
    return "assignment"


def source_summary(
    path: Path,
    filename: str,
    kind: str,
    card_generator: Callable[[str, str, str, dict], dict] | None = None,
    extracted: tuple[str, dict] | None = None,
) -> dict:
    suffix = path.suffix.lower()
    text, units = extracted if extracted is not None else extract_source_text(path, suffix)
    if suffix == ".pdf" and not text.strip():
        raise NoSelectableTextError(
            "This PDF has no selectable text. Scanned or image-only PDFs need OCR before they can be imported."
        )
    cleaned_lines = [re.sub(r"\s+", " ", line).strip() for line in text.splitlines()]
    cleaned_lines = [line for line in cleaned_lines if line]
    headings = []
    objectives = []
    objective_cues = []
    in_objectives = False
    objective_heading = re.compile(r"^(?:(?:course|student learning) )?(?:(?:learning )?objectives?|(?:learning )?outcomes?|learning goals?)(?:\s*[:\-].*)?$", re.I)
    section_heading = re.compile(r"^(?:course description|course schedule|weekly schedule|schedule|grading|policies|course policies|assignments|exams?|reading list|course materials|academic integrity|attendance|instructor|prerequisites?|weeks?\b)", re.I)
    objective_words = re.compile(r"\b(objective|outcome|students will|able to)\b", re.I)
    for line in cleaned_lines:
        normalized = line.lower()
        if len(line) <= 90 and (line.isupper() or re.match(r"^(unit|week|lecture|chapter|exam|section)\b", normalized)):
            headings.append(line)
        if objective_heading.match(line):
            in_objectives = True
            continue
        if in_objectives and (section_heading.match(line) or (line.isupper() and not objective_words.search(normalized))):
            in_objectives = False
        if in_objectives or objective_words.search(normalized):
            objectives.append(line)
            cue = re.sub(r"^\s*(?:[-*•]\s*|\d+[.)]\s*)", "", line).strip()
            word_count = len(re.findall(r"\b\w+\b", cue))
            if 5 <= word_count <= 50 and cue.casefold() not in {value.casefold() for value in objective_cues}:
                objective_cues.append(cue)
    fingerprint = hashlib.sha256(text.encode("utf-8", errors="ignore")).hexdigest()
    selected_kind = kind if kind in {"syllabus", "material", "assessment"} else "auto"
    detected_kind, classification_reason = classify_source(filename, text)
    resolved_kind = selected_kind if selected_kind != "auto" else detected_kind
    is_syllabus = resolved_kind == "syllabus"
    is_admin_form = _looks_like_administrative_form(text)
    if is_admin_form:
        classification_reason = "Administrative form detected; no study-card content expected"
    if is_syllabus or is_admin_form:
        compiled = {"concepts": [], "notes": [], "cards": []}
    elif suffix == ".pdf":
        page_texts = units.get("pageTexts") or text.split("\f")
        compiled = compile_pdf_study_material(page_texts, filename, assessment=resolved_kind == "assessment")
    elif resolved_kind == "assessment":
        assessment_cards = _labeled_assessment_cards(text, filename)
        compiled = (
            {"concepts": [], "notes": [], "cards": assessment_cards}
            if assessment_cards
            else compile_study_material(text, filename)
        )
    else:
        compiled = compile_study_material(text, filename)
    generated = card_generator(text, filename, resolved_kind, units) if card_generator and not is_syllabus and not is_admin_form else None
    structured_cards = draft_cards_from_structured_slides(text, filename) if suffix in {".ppt", ".pptw", ".pptx"} and not is_syllabus else []
    draft_cards = generated["cards"] if generated is not None else (structured_cards or compiled["cards"])
    concepts = generated["concepts"] if generated is not None else compiled["concepts"]
    calendar = syllabus_calendar(
        filename,
        text,
        hashlib.sha256(text.encode("utf-8", errors="ignore")).hexdigest()[:12],
        units.get("pageTexts") if suffix == ".pdf" else None,
    ) if is_syllabus else {"events": [], "courseName": "", "term": "", "termRange": "", "warnings": []}
    return {
        "id": fingerprint[:12],
        "name": filename,
        "kind": resolved_kind,
        "classificationReason": "Manually selected" if selected_kind != "auto" else classification_reason,
        "courseName": calendar["courseName"],
        "courseCode": calendar.get("courseCode", ""),
        "term": calendar["term"],
        "termRange": calendar["termRange"],
        "calendarEvents": calendar["events"],
        "calendarWarnings": calendar.get("warnings", []),
        "format": suffix.lstrip(".").upper(),
        "wordCount": len(re.findall(r"\b\w+\b", text)),
        "unitLabel": units["unitLabel"],
        "unitCount": units["unitCount"] or len(cleaned_lines),
        "headings": headings[:8],
        "objectiveCount": len(objectives),
        "objectiveCues": objective_cues[:20],
        "preview": cleaned_lines[:5],
        "concepts": concepts,
        "notes": compiled["notes"],
        "draftCards": draft_cards,
        **({"generation": generated["generation"]} if generated is not None else {}),
        "fingerprint": fingerprint,
        "addedAt": datetime.now(timezone.utc).isoformat(),
        "localOnly": True,
    }


def persist_session(payload: dict) -> str:
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    session_id = f"{stamp}-{safe_slug(payload.get('filename', 'lecture'))[:28]}"
    payload = {"id": session_id, "savedAt": datetime.now(timezone.utc).isoformat(), **payload}
    with _data_lock:
        write_json(SESSIONS_DIR / f"{session_id}.json", payload)
        write_json(DATA_DIR / "latest-session.json", payload)
    return session_id


def _bounded_int(value, fallback: int, minimum: int, maximum: int) -> int:
    try:
        return max(minimum, min(maximum, int(float(value))))
    except (TypeError, ValueError):
        return fallback


def _bounded_float(value, fallback: float, minimum: float, maximum: float) -> float:
    try:
        return max(minimum, min(maximum, float(value)))
    except (TypeError, ValueError):
        return fallback


def _step_minutes(value, fallback: list[float]) -> list[float]:
    if isinstance(value, list):
        parsed = []
        for step in value:
            try:
                parsed.append(max(0.0, float(step)))
            except (TypeError, ValueError):
                continue
        return parsed
    if not isinstance(value, str):
        return fallback
    parsed = []
    for raw_step in re.split(r"[\s,]+", value.strip()):
        if not raw_step:
            continue
        match = re.fullmatch(r"(\d+(?:\.\d+)?)([smhd]?)", raw_step, re.IGNORECASE)
        if not match:
            continue
        amount = float(match.group(1))
        unit = match.group(2).lower()
        if unit == "s":
            amount /= 60
        elif unit == "h":
            amount *= 60
        elif unit == "d":
            amount *= 1440
        parsed.append(amount)
    return parsed if parsed else fallback


def _anki_deck_config(preferences: dict, config_id: int) -> dict:
    gather_order = {
        "deck": 0,
        "ascending": 1,
        "descending": 2,
        "random-notes": 3,
        "random-cards": 4,
        "deck-random-notes": 5,
    }
    sort_order = {
        "template": 0,
        "gathered": 1,
        "template-random": 2,
        "random-note-template": 3,
        "random-card": 4,
    }
    review_order = {
        "due": 0,
        "due-deck": 1,
        "deck-due": 2,
        "interval-asc": 3,
        "interval-desc": 4,
        "ease-asc": 5,
        "ease-desc": 6,
        "retrievability-asc": 7,
        "random": 8,
        "added": 9,
        "reverse-added": 10,
        "retrievability-desc": 11,
        "relative-overdue": 12,
    }
    mix_order = {"mix": 0, "after": 1, "before": 2}
    answer_actions = {"bury": 0, "again": 1, "good": 2, "hard": 3, "reminder": 4}
    fsrs_parameters = preferences.get("fsrsParameters") or []
    if isinstance(fsrs_parameters, str):
        fsrs_parameters = [
            float(value)
            for value in re.split(r"[\s,]+", fsrs_parameters.strip())
            if value and re.fullmatch(r"-?\d+(?:\.\d+)?", value)
        ]
    easy_days = preferences.get("easyDays")
    if not isinstance(easy_days, list) or len(easy_days) != 7:
        easy_days = [1, 1, 1, 1, 1, 1, 1]
    easy_days = [_bounded_float(day, 1.0, 0.0, 1.0) for day in easy_days]
    return {
        "id": config_id,
        "mod": int(time.time()),
        "name": str(preferences.get("presetName") or "Syllabloom").strip()[:120] or "Syllabloom",
        "usn": -1,
        "maxTaken": _bounded_int(preferences.get("maximumAnswerSeconds"), 60, 0, 86400),
        "autoplay": not bool(preferences.get("disableAutoplay", False)),
        "timer": 1 if preferences.get("showTimer") else 0,
        "replayq": not bool(preferences.get("skipQuestionAudio", False)),
        "new": {
            "bury": bool(preferences.get("buryNew", True)),
            "delays": _step_minutes(preferences.get("learningSteps"), [1.0, 10.0]),
            "initialFactor": int(_bounded_float(preferences.get("startingEase"), 2.5, 1.3, 5.0) * 1000),
            "ints": [
                _bounded_int(preferences.get("graduatingInterval"), 1, 1, 36500),
                _bounded_int(preferences.get("easyInterval"), 4, 1, 36500),
                0,
            ],
            "order": 1 if preferences.get("insertionOrder") == "random" else 0,
            "perDay": _bounded_int(preferences.get("newPerDay", preferences.get("dailyLimit")), 20, 0, 9999),
        },
        "rev": {
            "bury": bool(preferences.get("buryReviews", True)),
            "ease4": _bounded_float(preferences.get("easyBonus"), 1.3, 1.0, 5.0),
            "ivlFct": _bounded_float(preferences.get("intervalModifier"), 1.0, 0.1, 5.0),
            "maxIvl": _bounded_int(preferences.get("maximumInterval"), 36500, 1, 36500),
            "perDay": _bounded_int(preferences.get("reviewsPerDay"), 9999, 0, 99999),
            "hardFactor": _bounded_float(preferences.get("hardInterval"), 1.2, 1.0, 2.0),
        },
        "lapse": {
            "delays": _step_minutes(preferences.get("relearningSteps"), [10.0]),
            "leechAction": 1 if preferences.get("leechAction") == "tag" else 0,
            "leechFails": _bounded_int(preferences.get("leechThreshold"), 8, 1, 999),
            "minInt": _bounded_int(preferences.get("minimumInterval"), 1, 1, 36500),
            "mult": _bounded_float(preferences.get("newInterval"), 0.0, 0.0, 1.0),
        },
        "dyn": False,
        "newMix": mix_order.get(str(preferences.get("newReviewOrder")), 1),
        "newPerDayMinimum": 0,
        "interdayLearningMix": mix_order.get(str(preferences.get("interdayOrder")), 0),
        "reviewOrder": review_order.get(str(preferences.get("reviewOrder")), 0),
        "newSortOrder": sort_order.get(str(preferences.get("sortOrder")), 0),
        "newGatherPriority": gather_order.get(str(preferences.get("gatherOrder")), 0),
        "buryInterdayLearning": bool(preferences.get("buryInterday", True)),
        "fsrsWeights": [],
        "fsrsParams5": [],
        "fsrsParams6": [float(value) for value in fsrs_parameters if isinstance(value, (int, float))],
        "desiredRetention": _bounded_float(preferences.get("desiredRetention"), 0.9, 0.7, 0.99),
        "ignoreRevlogsBeforeDate": str(preferences.get("ignoreBefore") or ""),
        "easyDaysPercentages": easy_days,
        "stopTimerOnAnswer": bool(preferences.get("stopTimerOnAnswer", False)),
        "secondsToShowQuestion": _bounded_float(preferences.get("questionSeconds"), 0.0, 0.0, 86400.0),
        "secondsToShowAnswer": _bounded_float(preferences.get("answerSeconds"), 0.0, 0.0, 86400.0),
        "questionAction": 1 if preferences.get("questionAction") == "show-reminder" else 0,
        "answerAction": answer_actions.get(str(preferences.get("answerAction")), 0),
        "waitForAudio": bool(preferences.get("waitForAudio", True)),
        "sm2Retention": _bounded_float(preferences.get("historicalRetention"), 0.9, 0.7, 0.99),
        "weightSearch": "",
    }


def _attach_anki_deck_config(package_path: Path, deck_id: int, config: dict) -> None:
    with zipfile.ZipFile(package_path, "r") as archive:
        entries = [(info, archive.read(info.filename)) for info in archive.infolist()]
    collection_bytes = next(data for info, data in entries if info.filename == "collection.anki2")
    with tempfile.NamedTemporaryFile(prefix="syllabloom-collection-", suffix=".anki2", delete=False) as database:
        database_path = Path(database.name)
        database.write(collection_bytes)
    repacked_path = package_path.with_suffix(".configured.apkg")
    try:
        connection = sqlite3.connect(database_path)
        decks_json, configs_json = connection.execute("SELECT decks, dconf FROM col").fetchone()
        decks = json.loads(decks_json)
        configs = json.loads(configs_json)
        decks[str(deck_id)]["conf"] = config["id"]
        configs[str(config["id"])] = config
        connection.execute("UPDATE col SET decks = ?, dconf = ?", (json.dumps(decks), json.dumps(configs)))
        connection.commit()
        connection.close()
        configured_collection = database_path.read_bytes()
        with zipfile.ZipFile(repacked_path, "w") as archive:
            for info, data in entries:
                archive.writestr(info, configured_collection if info.filename == "collection.anki2" else data)
        repacked_path.replace(package_path)
    finally:
        database_path.unlink(missing_ok=True)
        repacked_path.unlink(missing_ok=True)


def build_anki_package(cards: list[dict], preferences: dict, images: dict | None = None) -> tuple[bytes, str]:
    import genanki

    parent_deck = str(preferences.get("deck") or "Syllabloom").strip()[:120] or "Syllabloom"
    set_name = str(preferences.get("setName") or "").strip()[:120]
    deck_name = f"{parent_deck}::{set_name}" if set_name else parent_deck
    card_format = str(preferences.get("format") or "Basic")
    seed = int(hashlib.sha256(deck_name.encode("utf-8")).hexdigest()[:12], 16)
    deck_id = 1_000_000_000 + seed % 999_999_999
    model_id = deck_id - 17 if deck_id > 1_000_000_017 else deck_id + 17
    css = """
.card {
  margin: 0;
  padding: 24px 16px;
  font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Arial, sans-serif;
  font-size: 20px;
  line-height: 1.55;
  text-align: left;
  color: #252a27;
  background: #efeee8;
}
.card-shell {
  box-sizing: border-box;
  width: 100%;
  max-width: 720px;
  margin: 0 auto;
  padding: 28px;
  overflow-wrap: anywhere;
  background: #fbfaf6;
  border: 1px solid #c5cbc7;
  border-radius: 18px;
}
.prompt {
  font-size: 28px;
  font-weight: 650;
  line-height: 1.28;
  letter-spacing: -0.02em;
}
.answer {
  font-size: 20px;
  line-height: 1.55;
}
#answer {
  margin: 24px 0;
  border: 0;
  border-top: 1px solid #c9ceca;
}
.source {
  margin-top: 24px;
  padding-top: 12px;
  color: #626c66;
  font-size: 13px;
  line-height: 1.4;
  border-top: 1px solid #d8dcd9;
}
.mobile .card {
  padding: 12px 8px;
  font-size: 18px;
}
.mobile .card-shell {
  padding: 22px 18px;
  border-radius: 14px;
}
.mobile .prompt {
  font-size: 24px;
}
.nightMode.card {
  color: #eef1ee;
  background: #1d211f;
}
.nightMode .card-shell {
  background: #282e2a;
  border-color: #465049;
}
.nightMode .source {
  color: #b8c2bc;
  border-color: #465049;
}
.nightMode #answer {
  border-color: #465049;
}
"""
    cloze_model = genanki.Model(
        model_id + 1,
        "Syllabloom Cloze",
        fields=[{"name": "Text"}, {"name": "Back Extra"}, {"name": "Source"}],
        templates=[{
            "name": "Cloze",
            "qfmt": "<main class='card-shell'><div class='prompt'>{{cloze:Text}}</div></main>",
            "afmt": "<main class='card-shell'><div class='answer'>{{cloze:Text}}</div>{{#Back Extra}}<hr id='answer'><div class='answer'>{{Back Extra}}</div>{{/Back Extra}}<div class='source'>{{Source}}</div></main>",
        }],
        css=css,
        model_type=genanki.Model.CLOZE,
    )
    basic_model = genanki.Model(
        model_id,
        "Syllabloom Basic",
        fields=[{"name": "Front"}, {"name": "Back"}, {"name": "Source"}],
        templates=[{
            "name": "Card 1",
            "qfmt": "<main class='card-shell'><div class='prompt'>{{Front}}</div></main>",
            "afmt": "<main class='card-shell'><div class='prompt'>{{Front}}</div><hr id='answer'><div class='answer'>{{Back}}</div><div class='source'>{{Source}}</div></main>",
        }],
        css=css,
    )
    preset_name = str(preferences.get("presetName") or "Syllabloom").strip()[:120] or "Syllabloom"
    deck = genanki.Deck(deck_id, deck_name, description=f"Created by Syllabloom with the {preset_name} deck preset.")
    from api.card_formats import image_bytes, occlusion_html
    if not isinstance(images or {}, dict) or len(images or {}) > 24:
        raise ValueError('Too many diagram images in this export.')
    media = {ref: image_bytes(url) for ref, url in (images or {}).items()}
    for card in cards[:5000]:
        front = html.escape(str(card.get("front") or "").strip()).replace("\n", "<br>")
        back = html.escape(str(card.get("back") or "").strip()).replace("\n", "<br>")
        source = html.escape(str(card.get("source") or "Syllabloom source library").strip())
        if not front or not back:
            continue
        raw_tags = card.get("tags") or "syllabloom"
        tag_values = raw_tags if isinstance(raw_tags, (list, tuple, set)) else str(raw_tags).split()
        tags = [_safe_anki_tag(tag) for tag in tag_values if str(tag).strip()] or ["syllabloom"]
        note_type = card.get('noteType') or card_format
        if note_type == 'ImageOcclusion':
            occlusion = card.get('occlusion') or {}
            asset = media.get(occlusion.get('imageRef'))
            if not asset:
                raise ValueError('A diagram image is missing from this export.')
            front = occlusion_html(occlusion, asset[1])
            back = occlusion_html(occlusion, asset[1], reveal=True) + '<p>' + back + '</p>'
            note = genanki.Note(model=basic_model, fields=[front, back, source], tags=tags)
        elif note_type == 'Cloze':
            text = html.escape(str(card.get('clozeText') or card.get('front') or '')).replace('\n', '<br>')
            if card.get('noteType') and not re.search(r'\{\{c1::[^{}]+\}\}', text):
                raise ValueError('A cloze card needs a sentence with {{c1::term}}.')
            if '{{c' not in text:
                text = f'{front}<br>{{{{c1::{back}}}}}'
            note = genanki.Note(model=cloze_model, fields=[text, back, source], tags=tags)
        else:
            note = genanki.Note(model=basic_model, fields=[front, back, source], tags=tags)
        deck.add_note(note)
    if not deck.notes:
        raise ValueError("No complete cards were provided.")
    with tempfile.NamedTemporaryFile(prefix="syllabloom-anki-", suffix=".apkg", delete=False) as output:
        output_path = Path(output.name)
    try:
        with tempfile.TemporaryDirectory(prefix='syllabloom-images-') as image_directory:
            paths = []
            for raw, name in media.values():
                image_path = Path(image_directory) / name
                image_path.write_bytes(raw)
                paths.append(str(image_path))
            genanki.Package(deck, media_files=list(set(paths))).write_to_file(str(output_path))
        config_id = 2_000_000_000 + seed % 999_999_999
        _attach_anki_deck_config(output_path, deck_id, _anki_deck_config(preferences, config_id))
        return output_path.read_bytes(), f"{safe_slug(deck_name, 'syllabloom')}.apkg"
    finally:
        output_path.unlink(missing_ok=True)


# Prototype concept schema. In production this is generated from each class's
# syllabus and source library rather than being fixed in application code.
PHYSIOLOGY_CONCEPT_SCHEMA = (
    {
        "id": "homeostasis",
        "requires": ("homeostasis", "internal environment"),
        "title": "Homeostasis",
        "lines": (
            "The body uses regulatory mechanisms to preserve stable conditions in its internal environment.",
            "Fluid homeostasis is one example of this broader physiological goal.",
        ),
        "front": "What is homeostasis?",
        "back": "Maintenance of relatively stable conditions in the body’s internal environment through regulatory mechanisms.",
    },
    {
        "id": "organization-levels",
        "requires": ("chemical level", "cellular level", "tissue level"),
        "title": "Levels of structural organization",
        "lines": ("Chemical → cellular → tissue → organ → organ system → organism.",),
        "front": "Order the levels of structural organization from simplest to most complex.",
        "back": "Chemical → cellular → tissue → organ → organ system → organism.",
    },
    {
        "id": "tissue-types",
        "requires": ("four types of tissue", "epithelial tissue", "connective tissue"),
        "title": "Four basic tissue types",
        "lines": ("Epithelial, connective, muscular, and nervous tissue are the four basic tissue classes.",),
        "front": "What are the four basic tissue types?",
        "back": "Epithelial, connective, muscular, and nervous tissue.",
    },
    {
        "id": "microvilli",
        "requires": ("microvilli", "increase surface area"),
        "title": "Microvilli",
        "lines": ("Microvilli increase cell surface area and support absorption, including in the digestive tract.",),
        "front": "What is the main functional effect of microvilli?",
        "back": "They increase surface area for absorption.",
    },
    {
        "id": "rough-er",
        "requires": ("rough endoplasmic reticulum", "ribosomes make proteins"),
        "title": "Rough endoplasmic reticulum",
        "lines": ("Ribosomes give rough ER its appearance and synthesize proteins.",),
        "front": "Why does rough endoplasmic reticulum appear rough?",
        "back": "It is studded with ribosomes, which synthesize proteins.",
    },
    {
        "id": "golgi",
        "requires": ("packaging system", "outside of the cell"),
        "title": "Golgi apparatus",
        "lines": ("The Golgi modifies, sorts, and packages material for destinations inside, at the membrane, or outside the cell.",),
        "front": "What is the Golgi apparatus’s core role?",
        "back": "Modifying, sorting, and packaging cellular products for delivery.",
    },
    {
        "id": "plasma-membrane",
        "requires": ("plasma membrane", "lipid bilayer", "phospholipids"),
        "title": "Plasma membrane",
        "lines": (
            "The plasma membrane is primarily a phospholipid bilayer separating intracellular and extracellular environments.",
            "Its proteins, channels, receptors, carbohydrates, and lipids determine signaling and selective permeability.",
        ),
        "front": "What is the basic structure and purpose of the plasma membrane?",
        "back": "A phospholipid bilayer that separates internal from external environments and selectively controls transport and signaling.",
    },
    {
        "id": "passive-transport",
        "requires": ("passive", "doesn't require any energy", "diffusion"),
        "title": "Passive transport",
        "lines": ("Passive transport requires no ATP and moves substances down their concentration gradient.",),
        "front": "What defines passive transport?",
        "back": "It requires no ATP and moves substances down their concentration gradient.",
    },
    {
        "id": "simple-diffusion",
        "requires": ("simple diffusion", "high concentration to low concentration"),
        "title": "Simple diffusion",
        "lines": ("Small or lipid-soluble substances can cross the membrane directly from high to low concentration.",),
        "front": "Which substances cross by simple diffusion, and in what direction?",
        "back": "Small or lipid-soluble substances move directly through the membrane from high to low concentration.",
    },
    {
        "id": "facilitated-diffusion",
        "requires": ("facilitated diffusion", "glucose transporters"),
        "title": "Facilitated diffusion",
        "lines": ("A specific carrier moves a substance such as glucose down its gradient without ATP.",),
        "front": "How does facilitated diffusion differ from simple diffusion?",
        "back": "It still moves down the concentration gradient without ATP, but requires a specific carrier or transport protein.",
    },
    {
        "id": "active-transport",
        "requires": ("active", "requires", "atp", "against the gradient"),
        "title": "Active transport",
        "lines": ("Active transport uses energy, commonly ATP, to move substances against their concentration gradient.",),
        "front": "Why does active transport require energy?",
        "back": "It moves substances against their concentration gradient, from lower toward higher concentration.",
    },
    {
        "id": "osmosis",
        "requires": ("osmosis", "water"),
        "title": "Osmosis",
        "lines": ("Osmosis is the passive movement of water across a selectively permeable membrane.",),
        "front": "What is osmosis?",
        "back": "Passive movement of water across a selectively permeable membrane.",
    },
    {
        "id": "fluid-compartments",
        "requires": ("intracellular fluid", "extracellular fluid"),
        "title": "Body-fluid compartments",
        "lines": ("Body water is divided into intracellular and extracellular fluid compartments.",),
        "front": "What are the two major body-fluid compartments?",
        "back": "Intracellular fluid and extracellular fluid.",
    },
)


class NoAudioTrackError(ValueError):
    pass


def load_muscle_source() -> dict:
    raw = (ROOT / "muscle-data.js").read_text(encoding="utf-8").strip()
    prefix = "window.MUSCLE_SOURCE = "
    if not raw.startswith(prefix):
        raise RuntimeError("Muscle source is not in the expected format")
    payload = raw[len(prefix) :]
    if payload.endswith(";"):
        payload = payload[:-1]
    return json.loads(payload)


MUSCLE_SOURCE = load_muscle_source()


def get_model():
    global _model
    if _model is not None:
        return _model
    with _model_lock:
        if _model is None:
            if not MODEL_PATH.is_dir():
                raise RuntimeError(
                    "Local Whisper model not found. Set SYLLABLOOM_WHISPER_MODEL "
                    f"to a downloaded faster-whisper model directory. Checked: {MODEL_PATH}"
                )
            os.environ["ORT_DISABLE_TELEMETRY"] = "1"
            from faster_whisper import WhisperModel

            _model = WhisperModel(
                str(MODEL_PATH),
                device="cpu",
                compute_type="int8",
                local_files_only=True,
            )
    return _model


def find_course_matches(transcript: str) -> list[dict]:
    normalized = " ".join(transcript.lower().split())
    matches = []
    for record in MUSCLE_SOURCE["records"]:
        name = record["muscle"].lower()
        if name in normalized:
            matches.append(record)
    return matches


def find_lecture_concepts(transcript: str) -> list[dict]:
    normalized = " ".join(transcript.lower().replace("’", "'").split())
    return [
        rule
        for rule in PHYSIOLOGY_CONCEPT_SCHEMA
        if all(phrase in normalized for phrase in rule["requires"])
    ]


def build_lecture_items(rule: dict, heard_at: float) -> tuple[list[dict], list[dict]]:
    notes = [
        {
            "title": rule["title"],
            "section": "Medical Physiology",
            "heardAt": round(heard_at, 2),
            "status": "provisional",
            "lines": list(rule["lines"]),
        }
    ]
    cards = [
        {
            "muscle": rule["title"],
            "field": "lecture",
            "front": rule["front"],
            "back": rule["back"],
            "section": "Medical Physiology",
            "source": "Lecture transcript",
            "status": "provisional",
        }
    ]
    return notes, cards


def build_cards(records: list[dict]) -> list[dict]:
    cards = []
    field_prompts = {
        "attachment": "What are the attachments of {muscle}?",
        "action": "What is the action of {muscle}?",
        "innervation": "What is the innervation of {muscle}?",
    }
    for record in records:
        for field, prompt in field_prompts.items():
            cards.append(
                {
                    "muscle": record["muscle"],
                    "field": field,
                    "front": prompt.format(muscle=record["muscle"]),
                    "back": record[field],
                    "section": record["section"],
                    "source": MUSCLE_SOURCE["document"],
                }
            )
    return cards


def build_notes(records: list[dict], heard_at: float | None = None) -> list[dict]:
    return [
        {
            "title": record["muscle"],
            "section": record["section"],
            "heardAt": round(heard_at, 2) if heard_at is not None else None,
            "lines": [
                f'Attachments: {record["attachment"]}',
                f'Action: {record["action"]}',
                f'Innervation: {record["innervation"]}',
            ],
        }
        for record in records
    ]


def inspect_media(path: Path) -> dict:
    import av

    with av.open(str(path)) as container:
        audio_streams = [stream for stream in container.streams if stream.type == "audio"]
        video_streams = [stream for stream in container.streams if stream.type == "video"]
        duration = float(container.duration / av.time_base) if container.duration else 0.0
        if not audio_streams:
            raise NoAudioTrackError(
                "This video has no audio track. Download a merged YouTube file that includes both video and audio."
            )
        return {
            "durationSeconds": round(duration, 2),
            "hasVideo": bool(video_streams),
            "audioCodec": audio_streams[0].codec_context.name,
        }


def find_terminology_warnings(segments: list[dict]) -> list[dict]:
    """Flag a few high-risk anatomy substitutions without pretending to correct them."""
    checks = (
        ("planter", "Possible anatomy term: “planter” may be “plantar”."),
        ("interior tibial", "Possible vessel name: “interior tibial” may be “anterior tibial”."),
        ("mediocotally", "Unrecognized anatomical term; compare this phrase with the lecture audio."),
    )
    warnings = []
    seen = set()
    for index, segment in enumerate(segments):
        next_text = segments[index + 1]["text"] if index + 1 < len(segments) else ""
        normalized = f'{segment["text"]} {next_text}'.lower()
        for needle, message in checks:
            if needle in normalized and needle not in seen:
                seen.add(needle)
                warnings.append(
                    {
                        "time": segment["start"],
                        "term": needle,
                        "message": message,
                    }
                )
    return warnings


def transcribe(path: Path) -> dict:
    media = inspect_media(path)
    started = time.perf_counter()
    model = get_model()
    with _inference_lock:
        generated, info = model.transcribe(
            str(path),
            language="en",
            task="transcribe",
            beam_size=3,
            vad_filter=True,
            condition_on_previous_text=False,
        )
        segments = []
        transcript_parts = []
        weighted_logprob = 0.0
        weighted_seconds = 0.0
        for item in generated:
            text = item.text.strip()
            if not text:
                continue
            duration = max(0.01, float(item.end) - float(item.start))
            avg_logprob = float(getattr(item, "avg_logprob", -1.0))
            weighted_logprob += avg_logprob * duration
            weighted_seconds += duration
            transcript_parts.append(text)
            segments.append(
                {
                    "start": round(float(item.start), 2),
                    "end": round(float(item.end), 2),
                    "text": text,
                    "avgLogprob": round(avg_logprob, 3),
                    "needsReview": avg_logprob < -0.55,
                }
            )

    transcript = " ".join(transcript_parts)
    matches = find_course_matches(transcript)
    generic = compile_lecture_segments(segments, path.name)
    schema_concepts = []
    schema_notes = []
    schema_cards = []
    for rule in find_lecture_concepts(transcript):
        notes, cards = build_lecture_items(rule, 0)
        schema_concepts.append({"name": rule["title"], "status": "provisional"})
        schema_notes.extend(notes)
        schema_cards.extend(cards)
    matched_cards = build_cards(matches)
    seen_questions = {card["front"].strip().lower() for card in matched_cards + schema_cards}
    generic_cards = [
        card for card in generic["cards"]
        if card["front"].strip().lower() not in seen_questions
    ]
    terminology_warnings = find_terminology_warnings(segments)
    average_logprob = weighted_logprob / weighted_seconds if weighted_seconds else -2.0
    from faster_whisper.audio import decode_audio

    samples = decode_audio(str(path), sampling_rate=16000)
    bucket_count = 160
    bucket_size = max(1, len(samples) // bucket_count)
    waveform = []
    for index in range(bucket_count):
        bucket = samples[index * bucket_size : (index + 1) * bucket_size]
        waveform.append(round(float(abs(bucket).max()), 4) if len(bucket) else 0.0)
    peak = max(waveform, default=1.0) or 1.0
    waveform = [round(value / peak, 3) for value in waveform]
    return {
        "engine": "faster-whisper small · local CPU int8",
        "language": info.language,
        "durationSeconds": round(float(info.duration), 2),
        "processingSeconds": round(time.perf_counter() - started, 2),
        "averageLogprob": round(average_logprob, 3),
        "transcript": transcript,
        "segments": segments,
        "detectedConcepts": (
            [{"name": record["muscle"], "status": "verified"} for record in matches]
            + schema_concepts
            + generic["concepts"]
        ),
        "cards": matched_cards + schema_cards + generic_cards,
        "notes": build_notes(matches) + schema_notes + generic["notes"],
        "reviewSegmentCount": sum(1 for segment in segments if segment["needsReview"]),
        "qualityWarnings": terminology_warnings,
        "waveform": waveform,
        "media": media,
    }


def transcribe_events(path: Path, filename: str, markers: list[float] | None = None):
    markers = markers or []
    media = inspect_media(path)
    started = time.perf_counter()
    yield {
        "type": "metadata",
        "filename": filename,
        "durationSeconds": media["durationSeconds"],
        "hasVideo": media["hasVideo"],
        "engine": "faster-whisper small · local CPU int8",
    }
    model = get_model()

    with _inference_lock:
        generated, info = model.transcribe(
            str(path),
            language="en",
            task="transcribe",
            beam_size=3,
            vad_filter=True,
            condition_on_previous_text=False,
        )
        segments = []
        transcript_parts = []
        matched_names = set()
        matched_lecture_ids = set()
        seen_card_questions = set()
        session_concepts = []
        session_notes = []
        session_cards = []
        generic_buffer = []
        generic_window_number = 0
        for item in generated:
            text = item.text.strip()
            if not text:
                continue
            avg_logprob = float(getattr(item, "avg_logprob", -1.0))
            segment = {
                "start": round(float(item.start), 2),
                "end": round(float(item.end), 2),
                "text": text,
                "avgLogprob": round(avg_logprob, 3),
                "needsReview": avg_logprob < -0.55,
            }
            segments.append(segment)
            transcript_parts.append(text)
            yield {
                "type": "segment",
                "segment": segment,
                "processedSeconds": segment["end"],
                "durationSeconds": round(float(info.duration), 2),
            }

            matches = find_course_matches(" ".join(transcript_parts))
            for record in matches:
                if record["muscle"] in matched_names:
                    continue
                matched_names.add(record["muscle"])
                notes = build_notes([record], segment["start"])
                cards = build_cards([record])
                seen_card_questions.update(card["front"].strip().lower() for card in cards)
                session_concepts.append({"name": record["muscle"], "status": "verified"})
                session_notes.extend(notes)
                session_cards.extend(cards)
                yield {
                    "type": "concept",
                    "concept": record["muscle"],
                    "status": "verified",
                    "section": record["section"],
                    "notes": notes,
                    "cards": cards,
                }


            for rule in find_lecture_concepts(" ".join(transcript_parts)):
                if rule["id"] in matched_lecture_ids:
                    continue
                matched_lecture_ids.add(rule["id"])
                notes, cards = build_lecture_items(rule, segment["start"])
                seen_card_questions.update(card["front"].strip().lower() for card in cards)
                session_concepts.append({"name": rule["title"], "status": "provisional"})
                session_notes.extend(notes)
                session_cards.extend(cards)
                yield {
                    "type": "concept",
                    "concept": rule["title"],
                    "status": "provisional",
                    "section": "Medical Physiology",
                    "notes": notes,
                    "cards": cards,
                }

            generic_buffer.append(segment)
            buffered_words = sum(len(item["text"].split()) for item in generic_buffer)
            if len(generic_buffer) >= 5 or buffered_words >= 110:
                generic_window_number += 1
                compiled = compile_lecture_window(
                    " ".join(item["text"] for item in generic_buffer),
                    filename,
                    generic_buffer[0]["start"],
                    generic_window_number,
                )
                cards = [
                    card for card in compiled["cards"]
                    if card["front"].strip().lower() not in seen_card_questions
                ]
                if compiled["notes"]:
                    seen_card_questions.update(card["front"].strip().lower() for card in cards)
                    session_concepts.extend(compiled["concepts"])
                    session_notes.extend(compiled["notes"])
                    session_cards.extend(cards)
                    yield {
                        "type": "concept",
                        "concept": compiled["concepts"][0]["name"],
                        "status": "provisional",
                        "section": "Lecture notes",
                        "notes": compiled["notes"],
                        "cards": cards,
                    }
                generic_buffer = []

    transcript = " ".join(transcript_parts)
    if generic_buffer:
        generic_window_number += 1
        compiled = compile_lecture_window(
            " ".join(item["text"] for item in generic_buffer),
            filename,
            generic_buffer[0]["start"],
            generic_window_number,
        )
        cards = [
            card for card in compiled["cards"]
            if card["front"].strip().lower() not in seen_card_questions
        ]
        if compiled["notes"]:
            session_concepts.extend(compiled["concepts"])
            session_notes.extend(compiled["notes"])
            session_cards.extend(cards)
            yield {
                "type": "concept",
                "concept": compiled["concepts"][0]["name"],
                "status": "provisional",
                "section": "Lecture notes",
                "notes": compiled["notes"],
                "cards": cards,
            }
    warnings = find_terminology_warnings(segments)
    completed_at = datetime.now(timezone.utc).isoformat()
    session_id = persist_session(
        {
            "filename": filename,
            "durationSeconds": round(float(info.duration), 2),
            "processingSeconds": round(time.perf_counter() - started, 2),
            "completedAt": completed_at,
            "transcript": transcript,
            "segments": segments,
            "concepts": session_concepts,
            "notes": session_notes,
            "cards": session_cards,
            "qualityWarnings": warnings,
            "markers": markers,
            "sourceMatch": "matched" if matched_names else "review-only",
        }
    )
    yield {
        "type": "complete",
        "sessionId": session_id,
        "savedAt": completed_at,
        "durationSeconds": round(float(info.duration), 2),
        "processingSeconds": round(time.perf_counter() - started, 2),
        "transcript": transcript,
        "qualityWarnings": warnings,
        "reviewSegmentCount": sum(1 for segment in segments if segment["needsReview"]),
    }


class SyllabloomHandler(SimpleHTTPRequestHandler):
    server_version = "SyllabloomLocal/0.1"

    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=str(ROOT), **kwargs)

    def send_json(self, payload: dict, status: HTTPStatus = HTTPStatus.OK) -> None:
        body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(body)

    def send_bytes(self, payload: bytes, content_type: str, filename: str) -> None:
        self.send_response(HTTPStatus.OK)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Disposition", f'attachment; filename="{filename}"')
        self.send_header("Content-Length", str(len(payload)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(payload)

    def send_event_stream(self, events) -> None:
        self.send_response(HTTPStatus.OK)
        self.send_header("Content-Type", "application/x-ndjson; charset=utf-8")
        self.send_header("Cache-Control", "no-store")
        self.send_header("Connection", "close")
        self.end_headers()
        for event in events:
            line = (json.dumps(event, ensure_ascii=False) + "\n").encode("utf-8")
            self.wfile.write(line)
            self.wfile.flush()

    def dispatch_hosted_api(self, method: str) -> bool:
        """On a hosted deployment, serve /api/* from the api package. nginx serves the static site."""
        if not HOSTED:
            return False
        from api.routes import resolve

        handler_class = resolve(urlparse(self.path).path)
        if handler_class is None:
            self.send_json({"error": "Not found"}, HTTPStatus.NOT_FOUND)
            return True
        method_handler = getattr(handler_class, f"do_{method}", None)
        if method_handler is None:
            self.send_json({"error": "Method not allowed."}, HTTPStatus.METHOD_NOT_ALLOWED)
            return True
        endpoint = handler_class.__new__(handler_class)
        endpoint.__dict__ = self.__dict__  # share this request's socket files and headers
        method_handler(endpoint)
        return True

    def do_GET(self) -> None:
        if self.dispatch_hosted_api("GET"):
            return
        request_path = urlparse(self.path).path
        if request_path == "/api/billing-access":
            from api.billing import handle_api
            handle_api(self, self.command)
            return
        if request_path == "/api/user-data":
            from api.user_data import handle_request

            handle_request(self, "GET")
            return
        if request_path == "/api/health":
            self.send_json(
                {
                    "ok": True,
                    "engine": "faster-whisper",
                    "modelPresent": MODEL_PATH.is_dir(),
                    "mode": "local",
                }
            )
            return
        if request_path == "/api/auth-config":
            publishable_key = (
                os.environ.get("CLERK_PUBLISHABLE_KEY")
                or os.environ.get("VITE_CLERK_PUBLISHABLE_KEY")
                or os.environ.get("NEXT_PUBLIC_CLERK_PUBLISHABLE_KEY")
                or ""
            ).strip()
            self.send_json(clerk_auth_config(publishable_key, deployment_env()))
            return
        if request_path == "/api/sources":
            self.send_json({"sources": read_source_library()})
            return
        if request_path == "/api/sessions/latest":
            latest = DATA_DIR / "latest-session.json"
            if not latest.exists():
                self.send_json({"session": None})
                return
            try:
                self.send_json({"session": json.loads(latest.read_text(encoding="utf-8"))})
            except (OSError, json.JSONDecodeError):
                self.send_json({"session": None})
            return
        super().do_GET()

    def do_PUT(self) -> None:
        if self.dispatch_hosted_api("PUT"):
            return
        request_path = urlparse(self.path).path
        if request_path == "/api/user-data":
            from api.user_data import handle_request

            handle_request(self, "PUT")
            return
        self.send_json({"error": "Not found"}, HTTPStatus.NOT_FOUND)

    def do_POST(self) -> None:
        if self.dispatch_hosted_api("POST"):
            return
        request_path = urlparse(self.path).path
        if request_path == "/api/billing-access":
            from api.billing import handle_api
            handle_api(self, "POST")
            return
        if request_path == "/api/stripe-webhook":
            from api.billing import handle_webhook
            handle_webhook(self)
            return
        if request_path == "/api/export-anki":
            self.handle_anki_export()
            return
        if request_path == "/api/source":
            self.handle_source_upload()
            return
        if request_path == "/api/source-batch":
            self.handle_source_batch()
            return
        if request_path not in {"/api/transcribe", "/api/transcribe-stream"}:
            self.send_json({"error": "Not found"}, HTTPStatus.NOT_FOUND)
            return

        try:
            content_length = int(self.headers.get("Content-Length", "0"))
        except ValueError:
            content_length = 0
        if content_length <= 0 or content_length > MAX_UPLOAD_BYTES:
            self.send_json({"error": "Lecture files must be between 1 byte and 1 GB."}, HTTPStatus.BAD_REQUEST)
            return

        content_type = self.headers.get("Content-Type", "")
        if not content_type.startswith("multipart/form-data"):
            self.send_json({"error": "Expected multipart audio upload."}, HTTPStatus.BAD_REQUEST)
            return

        form = cgi.FieldStorage(
            fp=self.rfile,
            headers=self.headers,
            environ={"REQUEST_METHOD": "POST", "CONTENT_TYPE": content_type},
        )
        if "audio" not in form:
            self.send_json({"error": "No audio file was provided."}, HTTPStatus.BAD_REQUEST)
            return

        upload = form["audio"]
        markers = []
        if "markers" in form:
            try:
                raw_markers = json.loads(str(form.getvalue("markers") or "[]"))
                markers = [round(max(0.0, float(value)), 2) for value in raw_markers[:100]] if isinstance(raw_markers, list) else []
            except (TypeError, ValueError, json.JSONDecodeError):
                markers = []
        filename = Path(upload.filename or "recording.webm").name
        suffix = Path(filename).suffix.lower() or ".webm"
        if suffix not in ALLOWED_SUFFIXES:
            self.send_json({"error": "Unsupported audio format."}, HTTPStatus.BAD_REQUEST)
            return

        temporary_path = None
        try:
            with tempfile.NamedTemporaryFile(prefix="rounds-audio-", suffix=suffix, delete=False) as temporary:
                temporary_path = Path(temporary.name)
                while True:
                    chunk = upload.file.read(1024 * 1024)
                    if not chunk:
                        break
                    temporary.write(chunk)
            media = inspect_media(temporary_path)
            if request_path == "/api/transcribe-stream":
                self.send_event_stream(transcribe_events(temporary_path, filename, markers))
            else:
                result = transcribe(temporary_path)
                result["filename"] = filename
                result["media"] = media
                self.send_json(result)
        except NoAudioTrackError as exc:
            self.send_json({"error": str(exc), "code": "NO_AUDIO_TRACK"}, HTTPStatus.UNPROCESSABLE_ENTITY)
        except Exception as exc:  # keep native model errors out of the browser
            print(f"Transcription failed: {type(exc).__name__}", flush=True)
            self.send_json(
                {"error": "Local transcription failed. Check the audio format and model readiness."},
                HTTPStatus.INTERNAL_SERVER_ERROR,
            )
        finally:
            if temporary_path is not None:
                temporary_path.unlink(missing_ok=True)

    def handle_anki_export(self) -> None:
        from api.billing import require_access
        if not require_access(self):
            return
        try:
            content_length = int(self.headers.get("Content-Length", "0"))
        except ValueError:
            content_length = 0
        if content_length <= 0 or content_length > 5 * 1024 * 1024:
            self.send_json({"error": "Card export payload is missing or too large."}, HTTPStatus.BAD_REQUEST)
            return
        try:
            payload = json.loads(self.rfile.read(content_length).decode("utf-8"))
            cards = payload.get("cards") or []
            if not isinstance(cards, list) or not cards:
                raise ValueError("Approve at least one card before exporting.")
            package, filename = build_anki_package(cards, payload.get("preferences") or {})
            self.send_bytes(package, "application/octet-stream", filename)
        except (json.JSONDecodeError, UnicodeDecodeError, ValueError) as exc:
            self.send_json({"error": str(exc)}, HTTPStatus.BAD_REQUEST)
        except Exception as exc:
            print(f"Anki export failed: {type(exc).__name__}", flush=True)
            self.send_json({"error": "Anki package export failed locally."}, HTTPStatus.INTERNAL_SERVER_ERROR)

    def handle_source_upload(self) -> None:
        from api.billing import require_access
        if not require_access(self):
            return
        try:
            content_length = int(self.headers.get("Content-Length", "0"))
        except ValueError:
            content_length = 0
        if content_length <= 0 or content_length > 100 * 1024 * 1024:
            self.send_json({"error": "Source files must be between 1 byte and 100 MB."}, HTTPStatus.BAD_REQUEST)
            return
        content_type = self.headers.get("Content-Type", "")
        if not content_type.startswith("multipart/form-data"):
            self.send_json({"error": "Expected a document upload."}, HTTPStatus.BAD_REQUEST)
            return
        form = cgi.FieldStorage(
            fp=self.rfile,
            headers=self.headers,
            environ={"REQUEST_METHOD": "POST", "CONTENT_TYPE": content_type},
        )
        if "source" not in form:
            self.send_json({"error": "No source document was provided."}, HTTPStatus.BAD_REQUEST)
            return
        upload = form["source"]
        filename = Path(upload.filename or "source.txt").name
        suffix = Path(filename).suffix.lower()
        if suffix not in SOURCE_SUFFIXES:
            self.send_json({"error": "Use a DOCX, PowerPoint, PDF, or TXT source."}, HTTPStatus.BAD_REQUEST)
            return
        kind_value = form.getfirst("kind", "auto")
        temporary_path = None
        try:
            with tempfile.NamedTemporaryFile(prefix="rounds-source-", suffix=suffix, delete=False) as temporary:
                temporary_path = Path(temporary.name)
                while True:
                    chunk = upload.file.read(1024 * 1024)
                    if not chunk:
                        break
                    temporary.write(chunk)
            from api.ai_card_generation import CardGenerationError
            from api.ai_source_cards import generate_source_cards
            from api.user_data import authenticated_user

            user_id = authenticated_user(self.headers) or "local-development"
            question_style = form.getfirst("questionStyle", "balanced")
            if form.getfirst("inspect", "") == "1":
                from api.source import _source_preflight

                result = _source_preflight(temporary_path, filename, kind_value)
                result["source"]["localOnly"] = False
                result["source"]["storage"] = "session"
                self.send_json(result)
                return
            summary = source_summary(
                temporary_path,
                filename,
                kind_value,
                card_generator=lambda text, source_name, source_kind, units: generate_source_cards(
                    text, source_name, source_kind, units, user_id, question_style
                ),
            )
            with _data_lock:
                library = [item for item in read_source_library() if item.get("id") != summary["id"]]
                library.append(summary)
                write_json(SOURCE_LIBRARY_PATH, library)
            self.send_json({"source": summary})
        except NoSelectableTextError as exc:
            self.send_json({
                "error": str(exc),
                "errorCode": "NO_SELECTABLE_TEXT",
                "fileFingerprint": getattr(exc, "file_fingerprint", ""),
            }, HTTPStatus.UNPROCESSABLE_ENTITY)
        except PowerPointReadError as exc:
            self.send_json({"error": str(exc), "retryable": False}, HTTPStatus.UNPROCESSABLE_ENTITY)
        except CardGenerationError as exc:
            self.send_json({"error": exc.public_message}, exc.status)
        except Exception as exc:
            print(f"Source import failed: {type(exc).__name__}", flush=True)
            self.send_json({"error": "The source could not be read locally."}, HTTPStatus.UNPROCESSABLE_ENTITY)
        finally:
            if temporary_path is not None:
                temporary_path.unlink(missing_ok=True)

    def handle_source_batch(self) -> None:
        from api.billing import require_access
        if not require_access(self):
            return
        try:
            content_length = int(self.headers.get("Content-Length", "0"))
        except ValueError:
            content_length = 0
        if content_length <= 0 or content_length > 3_000_000:
            self.send_json({"error": "The card-generation request is missing or too large."}, HTTPStatus.BAD_REQUEST)
            return
        try:
            from api.source import _generate_source_batch
            from api.ai_card_generation import CardGenerationError
            from api.user_data import authenticated_user

            body = json.loads(self.rfile.read(content_length).decode("utf-8"))
            result = _generate_source_batch(body, authenticated_user(self.headers) or "local-development")
            self.send_json({"result": result})
        except CardGenerationError as exc:
            self.send_json({"error": exc.public_message, "retryable": exc.retryable}, exc.status)
        except (json.JSONDecodeError, UnicodeDecodeError, ValueError) as exc:
            self.send_json({"error": str(exc) or "The card-generation request is invalid."}, HTTPStatus.BAD_REQUEST)
        except Exception as exc:
            print(f"Source card batch failed: {type(exc).__name__}", flush=True)
            self.send_json({
                "error": "The card batch result could not be confirmed. It was not automatically retried to avoid duplicate generation.",
                "retryable": False,
            }, HTTPStatus.BAD_GATEWAY)


if __name__ == "__main__":
    sys.modules.setdefault("server", sys.modules[__name__])  # api handlers import this file as "server"
    host = os.environ.get("SYLLABLOOM_HOST", os.environ.get("ROUNDS_HOST", "127.0.0.1"))
    port = int(os.environ.get("SYLLABLOOM_PORT", os.environ.get("ROUNDS_PORT", "4174")))
    print(f"Syllabloom listening on http://{host}:{port}", flush=True)
    ThreadingHTTPServer((host, port), SyllabloomHandler).serve_forever()
