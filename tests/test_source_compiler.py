from __future__ import annotations

import tempfile
import json
import threading
import unittest
import zipfile
from http.server import ThreadingHTTPServer
from pathlib import Path
from urllib.error import HTTPError
from urllib.request import Request, urlopen

from api.source import handler as source_handler
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from server import NoSelectableTextError, _refine_pdf_quiz_pair, classify_source, compile_lecture_window, compile_study_material, source_summary


def write_text_pdf(path: Path, pages: list[str]) -> None:
    writer = PdfWriter()
    font = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    })
    font_reference = writer._add_object(font)
    for page_text in pages:
        page = writer.add_blank_page(width=612, height=792)
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_reference}),
        })
        commands = [b"BT /F1 10 Tf 54 740 Td"]
        for line in page_text.splitlines():
            escaped = line.encode("ascii", errors="replace")
            escaped = escaped.replace(b"\\", b"\\\\").replace(b"(", b"\\(").replace(b")", b"\\)")
            commands.extend((b"(" + escaped + b") Tj", b"0 -14 Td"))
        commands.append(b"ET")
        content = DecodedStreamObject()
        content.set_data(b" ".join(commands))
        page[NameObject("/Contents")] = writer._add_object(content)
    with path.open("wb") as stream:
        writer.write(stream)


class SourceCompilerTests(unittest.TestCase):
    def test_quiz_refinements_make_self_contained_concept_questions(self) -> None:
        system_calls = _refine_pdf_quiz_pair(
            "Give examples of system calls. Why are system calls necessary? Explain.",
            "Some examples of system calls include the create and delete calls for when the operating system starts and ends a process or when a privileged instruction is used. System calls are necessary because we want certain processes to be double checked by the operating system to ensure that both the process and operating system is protected.",
        )
        tsr = _refine_pdf_quiz_pair(
            "Give examples of TSR processes. In what situations should you consider a process to be a TSR process instead of just a regular process.",
            "Examples of TSR processes would be antivirus scanners and calculators. A process is considered a TSR process if it is invoked so often that it resides in the computer memory / address space. A regular process is not stored in the computer memory / address space.",
        )
        shells = _refine_pdf_quiz_pair(
            "Give examples of Unix shells. In what directory of the Unix file system can you find the shell program? When the user's shell is invoked?",
            "Examples of Unix Shells are zsh, csh, ksh. You can find the shell program in the /bin directory of the Unix file system. The user's shell is invoked whenever a user enters a command. The Unix shell runs in user mode.",
        )
        cards = system_calls + tsr + shells
        fronts = {front for front, _ in cards}
        all_copy = " ".join(front + " " + back for front, back in cards).casefold()

        self.assertIn("Which system-call examples concern process control and privileged instructions?", fronts)
        self.assertIn("What are examples of terminate-and-stay-resident (TSR) processes?", fronts)
        self.assertIn("Which Unix shells are given as examples?", fronts)
        self.assertIn("Where are Unix shell programs located?", fronts)
        self.assertNotIn("in this quiz", all_copy)
        self.assertNotIn("according to this quiz", all_copy)

    def test_auto_detects_syllabus_and_builds_calendar_without_study_cards(self) -> None:
        content = """PSY 241 Fall 2026 Course Syllabus
Course | PSY 241: Memory, Attention, and Decision Making
Term | Fall 2026 | August 25 to December 10
Week | Dates | Topics | Preparation | Due
1 | Aug 25, 27 | Cognitive evidence; operational definitions | Reader 1 | Reading check 1
2 | Nov 24, 26 | Anchoring; framing | Reader 2; no class Nov 26 | Project data due
Date | Milestone
September 10 | Memory lab proposal
December 1 and 3 | Final project presentations
"""
        with tempfile.NamedTemporaryFile(suffix=".txt", delete=False, mode="w", encoding="utf-8") as handle:
            path = Path(handle.name)
            handle.write(content)
        try:
            summary = source_summary(path, "PSY241_Fall_2026_Syllabus.txt", "auto")

            self.assertEqual(summary["kind"], "syllabus")
            self.assertEqual(summary["courseName"], "PSY 241: Memory, Attention, and Decision Making")
            self.assertEqual(summary["term"], "Fall 2026")
            self.assertEqual(summary["termRange"], "August 25 to December 10")
            self.assertEqual(summary["draftCards"], [])
            self.assertEqual(summary["notes"], [])
            self.assertEqual(summary["concepts"], [])
            events = summary["calendarEvents"]
            self.assertEqual(len(events), 6)
            dates = {event["date"] for event in events}
            self.assertIn("2026-08-25", dates)
            self.assertIn("2026-08-27", dates)
            self.assertIn("2026-11-24", dates)
            self.assertNotIn("2026-11-26", dates)
            self.assertIn("2026-12-03", dates)
            self.assertTrue(all(event["sourceId"] == summary["id"] for event in events))
            self.assertEqual(next(event for event in events if event["title"] == "Memory lab proposal")["type"], "assignment")
            mapped_event = next(event for event in events if event["title"].startswith("Week 1:"))
            self.assertIn("Aug 25, 27", mapped_event["sourceText"])
            self.assertEqual(mapped_event["yearSource"], "course term")
            self.assertEqual(mapped_event["sourceTerm"], "Fall 2026")
        finally:
            path.unlink(missing_ok=True)

    def test_syllabus_without_schedule_dates_imports_in_auto_and_manual_modes(self) -> None:
        content = """COURSE NAME: Operating Systems Principles
COURSE NUMBER: CSCI340
COURSE DESCRIPTION: Principles of operating-system design and implementation.
PREREQUISITE: CSCI 220, 240, and 313
COURSE OUTLINE
Processes
Threads
CPU Scheduling
GRADING POLICY
Quizzes 15%; homework 15%; exams 70%.
"""
        with tempfile.NamedTemporaryFile(suffix=".txt", delete=False, mode="w", encoding="utf-8") as handle:
            path = Path(handle.name)
            handle.write(content)
        try:
            auto = source_summary(path, "CSCI340_FA26.txt", "auto")
            forced = source_summary(path, "CSCI340_course_document.txt", "syllabus")

            self.assertEqual(classify_source("course-document.txt", content)[0], "syllabus")
            self.assertEqual(auto["kind"], "syllabus")
            self.assertEqual(auto["courseName"], "Operating Systems Principles")
            self.assertEqual(auto["courseCode"], "CSCI340")
            self.assertEqual(auto["term"], "Fall 2026")
            self.assertEqual(auto["calendarEvents"], [])
            self.assertEqual(auto["draftCards"], [])
            self.assertEqual(forced["kind"], "syllabus")
            self.assertEqual(forced["courseName"], "Operating Systems Principles")
            self.assertEqual(forced["termRange"], "")
            self.assertEqual(forced["calendarEvents"], [])
        finally:
            path.unlink(missing_ok=True)

    def test_pdf_syllabus_schedule_with_spaced_columns_adds_calendar_dates(self) -> None:
        content = """Course | PSY 241: Memory and Attention
Term | Fall 2026 | August to December
Instructor | Dr. Rivera
Week    Dates       Topics
1       Aug 25-27   Cognitive evidence and operational definitions
2       Sep 1, 3    Memory and attention
3       Sep 8       Decision making
Date        Milestone
Oct 8       Midterm 1
"""
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "course-schedule.pdf"
            write_text_pdf(path, [content])
            summary = source_summary(path, path.name, "auto")

        self.assertEqual(summary["kind"], "syllabus")
        events = summary["calendarEvents"]
        self.assertEqual({event["date"] for event in events}, {
            "2026-08-25", "2026-08-26", "2026-08-27", "2026-09-01", "2026-09-03", "2026-09-08", "2026-10-08"
        })
        self.assertEqual(next(event for event in events if event["date"] == "2026-10-08")["type"], "exam")

    def test_calendar_does_not_borrow_unrelated_year_and_flags_ambiguous_schedule_rows(self) -> None:
        content = """Course Syllabus
Copyright 2026 Example Publisher
Course Description: An introduction to marine ecosystems.
Week | Dates | Topics
1 | Aug 25 | Estuaries and food webs
"""
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "course-syllabus.pdf"
            write_text_pdf(path, [content])
            summary = source_summary(path, path.name, "syllabus")

        self.assertEqual(summary["calendarEvents"], [])
        self.assertTrue(any("calendar year was not clear" in warning for warning in summary["calendarWarnings"]))

    def test_schedule_row_with_explicit_year_imports_without_term_guess(self) -> None:
        content = """COURSE SYLLABUS
Weekly Schedule
Week | Dates | Topics
1 | Aug 25, 2026 | Estuaries and food webs
"""
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "course-syllabus.pdf"
            write_text_pdf(path, [content])
            summary = source_summary(path, path.name, "syllabus")

        self.assertEqual([event["date"] for event in summary["calendarEvents"]], ["2026-08-25"])
        self.assertEqual(summary["calendarEvents"][0]["yearSource"], "schedule row")

    def test_same_day_milestones_are_preserved_and_no_class_does_not_hide_deadlines(self) -> None:
        content = """PSY 241 Fall 2026 Course Syllabus
Weekly Schedule
Week | Dates | Topics | Reading | Due
10 | Nov 24, 26 | Memory | Reader 10 | Project data due Nov 26
No class Nov 26
Date | Milestone
Nov 26 | Midterm Exam 1
Nov 26 | Midterm Exam 2
"""
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "psy241-syllabus.pdf"
            write_text_pdf(path, [content])
            summary = source_summary(path, path.name, "syllabus")

        same_day = [event for event in summary["calendarEvents"] if event["date"] == "2026-11-26"]
        self.assertEqual({event["title"] for event in same_day}, {"Project data", "Midterm Exam 1", "Midterm Exam 2"})
        self.assertEqual(len(same_day), 3)

    def test_auto_detect_handles_varied_syllabus_and_material_pdfs(self) -> None:
        fixtures = [
            (
                "download.pdf",
                """COURSE SYLLABUS
COURSE NAME: Biology of Plants
COURSE NUMBER: BIO220
INSTRUCTOR: Dr. Rivera
COURSE DESCRIPTION: Plant physiology from cellular energy through growth.
REQUIRED TEXT: Plant Physiology, Fourth Edition
LEARNING OBJECTIVES
Explain photosynthesis and water transport.
GRADING POLICY
Midterm 30%; final 35%; lab reports 35%.
COURSE SCHEDULE
Week 1: Plant cells
Week 2: Photosynthesis
""",
                "syllabus",
            ),
            (
                "csci340_course_outline.pdf",
                """COMPUTER SCIENCE 340 - OPERATING SYSTEMS PRINCIPLES
Instructor: Dr. Chen
Office Hours: Tuesday 2 pm
Course Description: Design of processes, memory management, and synchronization.
Prerequisites: CSCI 220 and CSCI 240
Course Goals: Explain scheduling and virtual memory.
Tentative Weekly Schedule
Week 1 Introduction
Week 2 Processes and threads
Grade Distribution: quizzes 15 percent, homework 15 percent, exams 70 percent.
""",
                "syllabus",
            ),
            (
                "course-policies.pdf",
                """COURSE INFORMATION
Course: HIST 205 - Migration and Modern Cities
Instructor: Professor Ali
Prerequisite: HIST 101
Required Readings: Weekly primary sources
Learning Outcomes: Evaluate historical evidence.
Grading: Essays 60 percent; seminar participation 40 percent.
Class Schedule: See the dated weekly reading plan.
""",
                "syllabus",
            ),
            (
                "hist215.pdf",
                """HIST 215 - Modern World History
Instructor: Professor Ali
Meeting Times: Monday and Wednesday at 11 am
Catalog Description: Migration, cities, and the modern world.
Required Texts: Weekly primary source readings
Learning Outcomes: Evaluate historical claims using evidence.
Evaluation Criteria: Essays 60 percent; seminar participation 40 percent.
Weekly Schedule: Topics and readings are listed by week.
""",
                "syllabus",
            ),
            (
                "plant-physiology-lecture.pdf",
                """Plant Physiology - Lecture 4
Photosynthesis converts light energy into chemical energy stored in sugars.
The light reactions occur in the thylakoid membrane and produce ATP and NADPH.
The Calvin cycle fixes carbon dioxide into carbohydrate in the chloroplast stroma.
Stomata regulate gas exchange and water loss in leaves.
""",
                "material",
            ),
            (
                "learning-theory-reading.pdf",
                """Operant Conditioning
Operant conditioning is learning in which consequences change the future probability of a behavior.
Positive reinforcement adds a desirable consequence after a behavior and increases that behavior.
Negative reinforcement removes an aversive consequence after a behavior and increases that behavior.
""",
                "material",
            ),
            (
                "practice-set-4.pdf",
                """Practice Exam 1 - CSCI 340
Question 1: What is the role of a process control block?
A. Store process state
B. Allocate a file name
Question 2: Which scheduler selects a process for the CPU?
A. Short-term scheduler
B. Page replacement
Question 3: What does a semaphore control?
A. Access to shared resources
B. Virtual address width
""",
                "assessment",
            ),
        ]
        with tempfile.TemporaryDirectory() as directory:
            for filename, content, expected_kind in fixtures:
                with self.subTest(filename=filename):
                    path = Path(directory) / filename
                    write_text_pdf(path, [content])
                    summary = source_summary(path, filename, "auto")
                    self.assertEqual(summary["kind"], expected_kind)
                    self.assertGreater(summary["wordCount"], 20)
                    if expected_kind == "syllabus":
                        self.assertEqual(summary["draftCards"], [])
                        forced = source_summary(path, filename, "syllabus")
                        self.assertEqual(forced["kind"], "syllabus")
                        self.assertEqual(forced["draftCards"], [])
                        if filename == "hist215.pdf":
                            self.assertEqual(summary["courseCode"], "HIST215")
                            self.assertEqual(summary["courseName"], "Modern World History")
                    elif expected_kind == "material":
                        self.assertTrue(summary["draftCards"])

            multi_page_path = Path(directory) / "course-overview.pdf"
            write_text_pdf(multi_page_path, [
                "CSCI340 - Operating Systems Principles\nInstructor: Dr. Chen\nCourse Description: Process and memory management.",
                "Prerequisites: CSCI 220\nGrading Policy: Exams and labs\nWeekly Schedule: Topics by date.",
            ])
            multi_page = source_summary(multi_page_path, "course-overview.pdf", "auto")
            self.assertEqual(multi_page["kind"], "syllabus")
            self.assertEqual(multi_page["courseCode"], "CSCI340")
            self.assertEqual(multi_page["courseName"], "Operating Systems Principles")

    def test_scanned_or_image_only_pdf_fails_with_a_clear_message(self) -> None:
        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as handle:
            path = Path(handle.name)
        try:
            writer = PdfWriter()
            writer.add_blank_page(width=612, height=792)
            with path.open("wb") as stream:
                writer.write(stream)
            with self.assertRaisesRegex(NoSelectableTextError, "no selectable text"):
                source_summary(path, "scanned-syllabus.pdf", "auto")
        finally:
            path.unlink(missing_ok=True)

    def test_source_upload_endpoint_honors_auto_and_manual_pdf_types(self) -> None:
        content = """COURSE SYLLABUS
COURSE NAME: Environmental Chemistry
COURSE NUMBER: CHEM310
TERM: Fall 2026
COURSE DESCRIPTION: Chemical processes in natural systems.
GRADING POLICY: Exams 60 percent; field work 40 percent.
Weekly Schedule
Week | Dates | Topics
1 | Oct 8 | Field sampling methods
Date | Milestone
Oct 15 | Midterm Exam
"""
        with tempfile.TemporaryDirectory() as directory:
            pdf_path = Path(directory) / "course.pdf"
            write_text_pdf(pdf_path, [content])
            pdf_bytes = pdf_path.read_bytes()
            api_server = ThreadingHTTPServer(("127.0.0.1", 0), source_handler)
            worker = threading.Thread(target=api_server.serve_forever, daemon=True)
            worker.start()
            try:
                def submit_pdf(selected_kind: str, filename: str, contents: bytes) -> tuple[int, dict]:
                    boundary = "----SyllabloomSourceImportTest"
                    body = b"".join((
                        f"--{boundary}\r\nContent-Disposition: form-data; name=\"kind\"\r\n\r\n{selected_kind}\r\n".encode(),
                        f"--{boundary}\r\nContent-Disposition: form-data; name=\"source\"; filename=\"{filename}\"\r\nContent-Type: application/pdf\r\n\r\n".encode(),
                        contents,
                        f"\r\n--{boundary}--\r\n".encode(),
                    ))
                    request = Request(
                        f"http://127.0.0.1:{api_server.server_port}/api/source",
                        data=body,
                        headers={"Content-Type": f"multipart/form-data; boundary={boundary}"},
                        method="POST",
                    )
                    try:
                        with urlopen(request, timeout=10) as response:
                            return response.status, json.loads(response.read())
                    except HTTPError as response:
                        return response.code, json.loads(response.read())

                for selected_kind in ("auto", "syllabus"):
                    status, payload = submit_pdf(selected_kind, "course.pdf", pdf_bytes)
                    self.assertEqual(status, 200)
                    self.assertEqual(payload["source"]["kind"], "syllabus")
                    self.assertEqual(payload["source"]["courseName"], "Environmental Chemistry")
                    event = next(item for item in payload["source"]["calendarEvents"] if item["type"] == "lecture")
                    self.assertEqual(event["date"], "2026-10-08")
                    self.assertIn("Field sampling methods", event["sourceText"])
                    self.assertEqual(event["sourceName"], "course.pdf")

                blank_path = Path(directory) / "blank.pdf"
                blank_writer = PdfWriter()
                blank_writer.add_blank_page(width=612, height=792)
                with blank_path.open("wb") as stream:
                    blank_writer.write(stream)
                status, payload = submit_pdf("auto", "scanned-syllabus.pdf", blank_path.read_bytes())
                self.assertEqual(status, 422)
                self.assertIn("no selectable text", payload["error"])
            finally:
                api_server.shutdown()
                api_server.server_close()
                worker.join(timeout=2)


    def test_auto_source_classification_distinguishes_assessments_from_material(self) -> None:
        assessment, assessment_reason = classify_source(
            "PSY241_Practice_Midterm.docx",
            "Question 1: Which memory system maintains active information? A. Sensory B. Working",
        )
        material, _ = classify_source(
            "PSY241_Week_3_Slides.pptx",
            "Working memory maintains and manipulates a limited amount of information for an active goal.",
        )

        self.assertEqual(assessment, "assessment")
        self.assertIn("file name", assessment_reason)
        self.assertEqual(material, "material")

    def test_auto_detection_does_not_call_numbered_forms_an_exam(self) -> None:
        state_form = """DOCUMENT AND CERTIFICATE COVER SHEET
New York State Department of State
Contact Information
1. Name: ______________________________
2. Mailing Address: ____________________
3. Telephone Number: ___________________
4. Signature: __________________________
"""
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "1239-f.pdf"
            write_text_pdf(path, [state_form])
            summary = source_summary(path, path.name, "auto")

        self.assertEqual(summary["kind"], "material")
        self.assertIn("Administrative form", summary["classificationReason"])
        self.assertEqual(summary["draftCards"], [])
        self.assertEqual(summary["notes"], [])

    def test_pdf_import_creates_notes_and_cards(self) -> None:
        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as handle:
            path = Path(handle.name)
        try:
            writer = PdfWriter()
            page = writer.add_blank_page(width=612, height=792)
            font = DictionaryObject({
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            })
            page[NameObject("/Resources")] = DictionaryObject({
                NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)}),
            })
            content = DecodedStreamObject()
            content.set_data(
                b"BT /F1 12 Tf 72 720 Td (Photosynthesis) Tj "
                b"0 -18 Td (Photosynthesis is the process plants use to convert light into chemical energy.) Tj "
                b"0 -18 Td (The light reactions occur in the thylakoid membrane.) Tj "
                b"0 -18 Td (The Calvin cycle is the pathway that fixes carbon dioxide into sugars.) Tj ET"
            )
            page[NameObject("/Contents")] = writer._add_object(content)
            with path.open("wb") as stream:
                writer.write(stream)

            summary = source_summary(path, "photosynthesis-reading.pdf", "material")
            self.assertGreaterEqual(len(summary["notes"]), 1)
            self.assertGreaterEqual(len(summary["draftCards"]), 2)
            self.assertTrue(any("photosynthesis" in card["front"].lower() for card in summary["draftCards"]))
        finally:
            path.unlink(missing_ok=True)

    def test_pdf_cards_reflow_wrapped_definitions_and_keep_page_sources(self) -> None:
        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as handle:
            path = Path(handle.name)
        try:
            write_text_pdf(path, ["\n".join([
                "CS 340",
                "Lecturer: Course Staff",
                "Read: Lecture notes",
                "Topics: Operating Systems",
                "Operating System: A collection of programs that act as intermediaries between the user and",
                "the computer hardware.",
                "Multiprogramming: one CPU, multiple processes ready for execution.",
                "",
                "Process States",
                "Ready (active) - the process is in main memory and available for execution.",
                "P 1 0",
                "C P U",
            ])])

            summary = source_summary(path, "operating-systems-notes.pdf", "material")
            cards = summary["draftCards"]
            fronts = {card["front"] for card in cards}

            self.assertIn("What is an operating system?", fronts)
            self.assertIn("What is multiprogramming?", fronts)
            self.assertTrue(all(card["source"].endswith("Page 1") for card in cards))
            self.assertFalse(any(card["front"].lower().startswith(("what is read?", "what is topics?")) for card in cards))
            os_card = next(card for card in cards if card["front"] == "What is an operating system?")
            self.assertIn("the computer hardware", os_card["back"])
            self.assertGreaterEqual(len(summary["notes"]), 1)
        finally:
            path.unlink(missing_ok=True)

    def test_answered_quiz_pdf_matches_numbered_questions_without_cross_section_bleed(self) -> None:
        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as handle:
            path = Path(handle.name)
        try:
            page_one = "\n".join([
                "Quiz 1",
                "1.1 What does a system call connect?",
                "1.2 Why is synchronization needed?",
                "Answers",
                "1.1 A system call is an interface between a user process and the operating system.",
                "1.2 Synchronization prevents a process from reading shared data while another process writes to it.",
            ])
            page_two = "\n".join([
                "2.1 What is the role of the command interpreter?",
                "2.2 What operating system is installed on my computer?",
                "2.3 When is command.com loaded for Unix?",
                "Answers",
                "2.1 The command interpreter is the interface between the user and the operating system.",
                "2.2 The OS installed on my computer is Windows.",
                "2.3 command.com is loaded into memory during boot time for Unix.",
            ])
            write_text_pdf(path, [page_one, page_two])

            summary = source_summary(path, "quiz-answers.pdf", "assessment")
            cards = summary["draftCards"]

            self.assertEqual({card["questionId"] for card in cards}, {"1.1", "1.2", "2.1"})
            self.assertEqual(len(cards), 3)
            self.assertTrue(all(card["source"].startswith("quiz-answers.pdf · Question ") for card in cards))
            self.assertTrue(all("What operating system is installed" not in card["back"] for card in cards))
            self.assertTrue(all("2.3 When is command.com" not in card["back"] for card in cards))
        finally:
            path.unlink(missing_ok=True)

    def test_pdf_cards_keep_comparisons_grammatical_and_io_answers_complete(self) -> None:
        content = """CS 340
Multiprogramming: one CPU, multiple processes ready for execution.
Multiprocessing: several processors (CPUs) are used on a single computer system to increase processing power.
Hard real-time system: guarantees that a critical task is done in time.
Soft real-time system: a critical real-time task gets priority over other tasks.
Interrupt Driven I/O
The CPU issues a command requesting I/O. Then the CPU continues to execute other instructions until the I/O module completes its work. At that point, the I/O module issues an interrupt.
"""
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "systems-lecture.pdf"
            write_text_pdf(path, [content])
            summary = source_summary(path, path.name, "material")

        cards = {card["front"]: card["back"] for card in summary["draftCards"]}
        comparison = cards["How does multiprogramming differ from multiprocessing?"]
        self.assertIn("uses one CPU, with multiple processes ready", comparison)
        self.assertIn("uses several processors (CPUs)", comparison)
        real_time = cards["How do hard and soft real-time systems differ?"]
        self.assertIn("guarantees that a critical task is completed in time", real_time)
        self.assertIn("gives a critical task priority", real_time)
        io_answer = cards["How does interrupt-driven I/O let the CPU keep working?"]
        self.assertIn("continues executing other instructions", io_answer)
        self.assertIn("interrupts the CPU when it finishes", io_answer)
        self.assertTrue(all(card["source"].endswith("Page 1") for card in summary["draftCards"]))

    def test_plain_text_import_creates_notes_and_cards(self) -> None:
        with tempfile.NamedTemporaryFile(suffix=".txt", delete=False, mode="w", encoding="utf-8") as handle:
            path = Path(handle.name)
            handle.write(
                "Operant conditioning\n"
                "Operant conditioning is learning in which consequences change behavior.\n"
                "Positive reinforcement is the addition of a desirable stimulus after a behavior.\n"
                "Negative reinforcement is the removal of an aversive stimulus after a behavior.\n"
            )
        try:
            summary = source_summary(path, "learning-theory.txt", "material")
            self.assertGreaterEqual(len(summary["notes"]), 1)
            self.assertGreaterEqual(len(summary["draftCards"]), 2)
            self.assertTrue(any("operant conditioning" in card["front"].lower() for card in summary["draftCards"]))
        finally:
            path.unlink(missing_ok=True)

    def test_word_import_creates_source_traced_cards(self) -> None:
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as handle:
            path = Path(handle.name)
        try:
            document = Document()
            document.add_heading("Supply and demand", level=1)
            document.add_paragraph("Demand is the quantity consumers are willing and able to buy at a given price.")
            document.add_paragraph("A price ceiling is a legal maximum price for a good or service.")
            document.add_paragraph("A binding price ceiling set below equilibrium creates a shortage.")
            document.save(path)

            summary = source_summary(path, "microeconomics-notes.docx", "material")
            self.assertGreaterEqual(len(summary["notes"]), 1)
            self.assertGreaterEqual(len(summary["draftCards"]), 2)
            self.assertTrue(all("microeconomics-notes.docx" in card["source"] for card in summary["draftCards"]))
        finally:
            path.unlink(missing_ok=True)
    def test_general_slides_create_concepts_notes_and_real_questions(self) -> None:
        text = """Slide 1
Cellular respiration
Cellular respiration is the process cells use to convert glucose into usable ATP.
Glycolysis occurs in the cytosol.
The citric acid cycle produces electron carriers for oxidative phosphorylation.
Slide 2
Action potentials
An action potential is a rapid change in membrane voltage.
Depolarization occurs when sodium enters the neuron.
Repolarization results from potassium leaving the neuron.
"""

        compiled = compile_study_material(text, "week-3-physiology.pptx")

        self.assertEqual(len(compiled["concepts"]), 2)
        self.assertEqual(len(compiled["notes"]), 2)
        self.assertGreaterEqual(len(compiled["cards"]), 4)
        fronts = {card["front"].lower() for card in compiled["cards"]}
        self.assertIn("what is cellular respiration?", fronts)
        self.assertIn("where does glycolysis occur?", fronts)
        self.assertIn("what is an action potential?", fronts)
        self.assertTrue(all("Slide" in card["source"] for card in compiled["cards"]))
        self.assertTrue(all(card["status"] == "verified" for card in compiled["cards"]))

    def test_psychology_slides_keep_only_direct_source_supported_cards(self) -> None:
        text = """Slide 1
Memory, Attention, and Decision Making
PSY 241 Unit lecture and cumulative review
How people hold information, lose it, find it again, and use it to choose
PSY 241
Slide 2
Questions for this unit
Which memory system is required by a task?
Identify one condition where automatic processing helps.
PSY 241
Slide 3
Memory systems differ by function and duration
Sensory memory briefly preserves modality-specific detail after stimulation ends.
Working memory maintains and manipulates a limited amount of information for an active goal.
Episodic memory represents personally experienced events with contextual detail.
PSY 241
Slide 4
Task switching creates measurable costs
A switch cost is the increase in response time or error when the required task changes from the previous trial.
Multitasking often means rapid switching.
PSY 241
Slide 5
Experimental design challenge
Question: Does retrieval practice improve delayed explanation more than restudy?
Independent variable: retrieval practice with feedback versus matched-time restudy.
Dependent variable: accuracy on a delayed explanation test one week later.
PSY 241
Slide 6
Key terms for review
Encoding specificity: cue effectiveness depends on overlap with encoding.
Consolidation: stabilization and reorganization after learning.
PSY 241
"""

        compiled = compile_study_material(text, "PSY241_Memory_Attention_Decision_Making.pptx")
        fronts = {card["front"] for card in compiled["cards"]}
        backs = [card["back"] for card in compiled["cards"]]
        concept_names = {concept["name"] for concept in compiled["concepts"]}

        self.assertIn("What does sensory memory briefly preserve?", fronts)
        self.assertIn("What is a switch cost?", fronts)
        self.assertIn("What does multitasking often mean?", fronts)
        self.assertIn("What outcome does the experiment test?", fronts)
        self.assertIn("What is the independent variable in the experiment?", fronts)
        self.assertIn("What is encoding specificity?", fronts)
        self.assertNotIn("Questions for this unit", concept_names)
        self.assertFalse(any(front.startswith("What are the key ideas") for front in fronts))
        self.assertFalse(any(front in {"What is question?", "What is independent variable?"} for front in fronts))
        self.assertTrue(all("PSY 241" not in back and "\n" not in back for back in backs))
        outcome_card = next(card for card in compiled["cards"] if card["front"] == "What outcome does the experiment test?")
        self.assertEqual(outcome_card["back"], "whether retrieval practice improves delayed explanation more than restudy")

    def test_short_author_and_relational_questions_keep_their_slide_context(self) -> None:
        text = """Slide 1
Capacity depends on organization
Cowan proposed a narrower focus of attention near four meaningful units under controlled conditions.
Slide 2
Type 1 and Type 2 processing
The distinction describes processing characteristics rather than fixed brain systems.
"""

        cards = compile_study_material(text, "PSY241_week-3.pptx")["cards"]
        fronts = {card["front"] for card in cards}

        self.assertIn('In “Capacity depends on organization,” what did Cowan propose?', fronts)
        self.assertIn('What does the distinction within “Type 1 and Type 2 processing” describe?', fronts)

    def test_card_templates_keep_grammar_and_answers_atomic(self) -> None:
        text = """Slide 1
Attention mechanisms
Source-monitoring errors occur when a person remembers a detail but misattributes its source.
Late-selection theories allow semantic processing before response selection.
Controlled processing requires attention, responds flexibly to goals, and becomes vulnerable when tasks compete.
Slide 2
Learning conditions
Interleaving mixes categories or problem types across practice.
Practice can reduce resource demands, but an automatic response can become costly when the mapping changes.
Slide 3
Bias
Search bias favors questions and evidence that could confirm an existing belief.
Interpretation bias treats ambiguous evidence as supportive of a preferred conclusion.
"""

        compiled = compile_study_material(text, "psychology-notes.pptx")
        cards = {card["front"]: card["back"] for card in compiled["cards"]}

        self.assertIn("When do source-monitoring errors occur?", cards)
        self.assertIn("What do late-selection theories allow?", cards)
        self.assertEqual(cards["What does controlled processing require?"], "attention")
        self.assertEqual(cards["What does interleaving mix?"], "categories or problem types across practice")
        self.assertIn("What does search bias favor?", cards)
        self.assertIn("What does interpretation bias treat?", cards)
        self.assertFalse(any(" practice can " in question.lower() for question in cards))

    def test_end_to_end_powerpoint_import_uses_normal_slide_text(self) -> None:
        presentation = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<p:presentation xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:sldIdLst><p:sldId id="256" r:id="rId1"/><p:sldId id="257" r:id="rId2"/></p:sldIdLst></p:presentation>"""
        relationships = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide" Target="slides/slide1.xml"/><Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide" Target="slides/slide2.xml"/></Relationships>"""

        def slide_xml(*lines: str) -> str:
            paragraphs = "".join(f"<a:p><a:r><a:t>{line}</a:t></a:r></a:p>" for line in lines)
            return f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<p:sld xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:cSld><p:spTree><p:sp><p:txBody><a:bodyPr/><a:lstStyle/>{paragraphs}</p:txBody></p:sp></p:spTree></p:cSld></p:sld>"""

        with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as handle:
            path = Path(handle.name)
        try:
            with zipfile.ZipFile(path, "w") as package:
                package.writestr("ppt/presentation.xml", presentation)
                package.writestr("ppt/_rels/presentation.xml.rels", relationships)
                package.writestr(
                    "ppt/slides/slide1.xml",
                    slide_xml(
                        "Renal filtration",
                        "Glomerular filtration is the movement of fluid from glomerular capillaries into Bowman's space.",
                        "Filtration rate depends on hydrostatic and oncotic pressures.",
                    ),
                )
                package.writestr(
                    "ppt/slides/slide2.xml",
                    slide_xml(
                        "Tubular reabsorption",
                        "Tubular reabsorption returns filtered substances to the blood.",
                        "Most sodium reabsorption occurs in the proximal tubule.",
                    ),
                )

            summary = source_summary(path, "renal-physiology.pptx", "material")
            self.assertEqual(summary["unitCount"], 2)
            self.assertEqual(len(summary["notes"]), 2)
            self.assertEqual(len(summary["concepts"]), 2)
            self.assertGreaterEqual(len(summary["draftCards"]), 2)
            self.assertTrue(any("glomerular filtration" in card["front"].lower() for card in summary["draftCards"]))
        finally:
            path.unlink(missing_ok=True)

    def test_lecture_window_creates_timestamped_notes_and_questions(self) -> None:
        compiled = compile_lecture_window(
            "An action potential is a rapid change in membrane voltage. Depolarization occurs when sodium enters the neuron. Potassium leaving the cell causes repolarization.",
            "Neurophysiology lecture.webm",
            125.4,
            2,
        )

        self.assertEqual(len(compiled["notes"]), 1)
        self.assertGreaterEqual(len(compiled["cards"]), 1)
        self.assertEqual(compiled["notes"][0]["heardAt"], 125.4)
        self.assertIn("02:05", compiled["cards"][0]["source"])
        self.assertEqual(compiled["cards"][0]["status"], "provisional")

    def test_lecture_compiler_rejects_promotional_and_broken_subjects(self) -> None:
        compiled = compile_lecture_window(
            "And is part of a group of extensor muscles. This muscle runs from the lateral tibia toward the ankle.",
            "lecture.webm",
            30,
            1,
        )
        self.assertFalse(any(card["front"].lower().startswith("what is and") for card in compiled["cards"]))

        promotional = compile_lecture_window(
            "This video is more fun than reading a textbook. Visit our site for more videos and interactive quizzes.",
            "lecture.webm",
            90,
            2,
        )
        self.assertEqual(promotional["cards"], [])
        self.assertEqual(promotional["notes"], [])

        transcript_fragments = compile_lecture_window(
            "My name is John Sullivan and this is the first lecture. "
            "The body to help maintain that homeostasis is all about vital mechanisms. "
            "The example he gave here is smooth muscle cells. "
            "Muscular and as well as nervous tissue those are four basic types. "
            "Instead what it does may rely primarily on something cocholine. "
            "Remember the intestines help with digestion. "
            "Could be is an entirely different function. "
            "You'll see that many different types of cells help catalyze products glucose.",
            "lecture.webm",
            1509,
            12,
        )
        self.assertEqual(transcript_fragments["concepts"], [])
        self.assertEqual(transcript_fragments["notes"], [])
        self.assertEqual(transcript_fragments["cards"], [])

    def test_generic_lecture_requires_a_clear_testable_fact(self) -> None:
        compiled = compile_lecture_window(
            "Inflation is a sustained increase in the general price level. "
            "Contractionary monetary policy decreases aggregate demand by raising interest rates.",
            "economics-lecture.m4a",
            420,
            4,
        )

        self.assertEqual(compiled["concepts"][0]["name"], "Inflation")
        self.assertEqual(len(compiled["cards"]), 2)
        self.assertEqual(compiled["cards"][0]["front"], "What is inflation?")
        self.assertTrue(all(card["status"] == "provisional" for card in compiled["cards"]))

    def test_lecture_compiler_rejects_singular_are_and_formats_unit_equivalence(self) -> None:
        compiled = compile_lecture_window(
            "Urine are typical again between 1 and 1.5 liters. "
            "Kilogram is equal to 2.2 pounds approximately.",
            "physiology-lecture.m4a",
            2522.37,
            106,
        )

        self.assertEqual(len(compiled["cards"]), 1)
        self.assertEqual(compiled["cards"][0]["front"], "What is one kilogram equal to?")
        self.assertEqual(compiled["cards"][0]["back"], "2.2 pounds approximately")
        self.assertEqual(compiled["concepts"][0]["name"], "Kilogram")

    def test_lecture_compiler_removes_discourse_words_from_inclusion_cards(self) -> None:
        compiled = compile_lecture_window(
            "Whole blood also contains cellular components, like red blood cells and white blood cells, right?",
            "physiology-lecture.m4a",
            3192.55,
            137,
        )

        self.assertEqual(compiled["concepts"][0]["name"], "Whole blood")
        self.assertEqual(compiled["cards"][0]["front"], "What does whole blood include?")
        self.assertEqual(
            compiled["cards"][0]["back"],
            "cellular components, like red blood cells and white blood cells",
        )

    def test_lecture_compiler_rejects_context_only_and_run_on_cards(self) -> None:
        fragments = (
            "The reverse is true with potassium. "
            "Milligram out of 100 ml is a 1% 10 milligrams out of 100 ml is 10% it's easy right now let's give you a sort of a real world example here. "
            "Millimoles is the number of sodium molecules in that. "
            "Osmoles are very important because we're gonna be talking about osmoles in the next slides. "
            "The idea is I want to know how many particles are in solution because that number directly affects fluid shift."
        )
        compiled = compile_lecture_window(
            fragments,
            "physiology-lecture.m4a",
            3600,
            150,
        )

        self.assertEqual(compiled["concepts"], [])
        self.assertEqual(compiled["notes"], [])
        self.assertEqual(compiled["cards"], [])

    def test_unrelated_subjects_create_specific_study_questions(self) -> None:
        subjects = {
            "chemistry.txt": (
                "Chemical kinetics\n"
                "Activation energy is the minimum energy required for a reaction to proceed.\n"
                "A catalyst decreases activation energy without being consumed by the reaction.\n",
                "What is activation energy?",
            ),
            "economics.txt": (
                "Monetary policy\n"
                "Inflation is a sustained increase in the general price level.\n"
                "Contractionary monetary policy decreases aggregate demand by raising interest rates.\n",
                "What is inflation?",
            ),
            "contract-law.txt": (
                "Contract formation\n"
                "Consideration is the exchange of value that supports an enforceable contract.\n"
                "An offer is a definite promise made with intent to be bound.\n",
                "What is consideration?",
            ),
            "cold-war-history.txt": (
                "Postwar Europe\n"
                "The Marshall Plan was a United States program that funded European economic recovery.\n"
                "The Berlin Airlift was an Allied operation that supplied West Berlin by air.\n",
                "What was the Marshall Plan?",
            ),
            "computer-science.txt": (
                "Data structures\n"
                "A hash table is a data structure that maps keys to storage locations.\n"
                "A collision is an event in which two keys map to the same location.\n",
                "What is a hash table?",
            ),
            "literature.txt": (
                "Narrative technique\n"
                "Dramatic irony is a technique in which the audience knows more than a character.\n"
                "A motif is a recurring image or idea that develops a theme.\n",
                "What is dramatic irony?",
            ),
        }

        for filename, (text, expected_question) in subjects.items():
            with self.subTest(filename=filename):
                compiled = compile_study_material(text, filename)
                self.assertGreaterEqual(len(compiled["notes"]), 1)
                self.assertGreaterEqual(len(compiled["cards"]), 2)
                self.assertIn(expected_question, {card["front"] for card in compiled["cards"]})
                self.assertTrue(all(card["source"] == filename for card in compiled["cards"]))

    def test_article_led_economics_and_neuroscience_material_generates_source_grounded_cards(self) -> None:
        economics = """Externalities and market failure
A negative externality imposes costs on third parties that are not reflected in the market price.
A positive externality creates benefits for third parties and can lead to underproduction.
A subsidy can encourage consumption or production when marginal social benefit exceeds marginal private benefit.
"""
        neuroscience = """Synaptic Transmission
Calcium entry triggers synaptic vesicles to release neurotransmitter into the synaptic cleft.
Neurotransmitters bind receptors on the postsynaptic membrane and change ion-channel activity.
Reuptake transporters return neurotransmitter from the synaptic cleft to the presynaptic neuron.
"""
        for filename, text in (("market-failure.txt", economics), ("synaptic-transmission.txt", neuroscience)):
            with self.subTest(filename=filename):
                compiled = compile_study_material(text, filename)
                cards = compiled["cards"]
                self.assertGreaterEqual(len(cards), 2)
                self.assertTrue(all(card["source"] == filename for card in cards))
                self.assertTrue(all(card["front"] and card["back"] for card in cards))
                self.assertFalse(any(card["front"].lower() == "what is answer?" for card in cards))

    def test_pdf_article_led_material_keeps_concise_but_complete_economics_facts(self) -> None:
        content = """Externalities and market failure
A negative externality imposes costs on third parties that are not reflected in the market price.
A positive externality creates benefits for third parties and can lead to underproduction.
A subsidy can encourage consumption or production when marginal social benefit exceeds marginal private benefit.
"""
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "externalities.pdf"
            write_text_pdf(path, [content])
            summary = source_summary(path, path.name, "material")

        cards = summary["draftCards"]
        self.assertGreaterEqual(len(cards), 3)
        self.assertTrue(all(card["source"].endswith("Page 1") for card in cards))
        self.assertTrue(any("externality" in card["front"].lower() for card in cards))
        self.assertIn("When can a subsidy encourage consumption or production?", {card["front"] for card in cards})

    def test_passive_science_statement_becomes_a_process_question_not_a_fake_definition(self) -> None:
        compiled = compile_study_material(
            "Cellular respiration\nAcetyl-CoA is oxidized in the mitochondrial matrix.",
            "cellular-respiration.txt",
        )
        card = compiled["cards"][0]
        self.assertEqual(card["front"], "What happens to Acetyl-CoA?")
        self.assertEqual(card["back"], "It is oxidized in the mitochondrial matrix")

    def test_plain_historical_narrative_generates_grounded_event_questions(self) -> None:
        text = """French Revolution: Causes and Turning Points
The Estates-General met in 1789 amid a fiscal crisis and disputes over political representation.
The National Assembly adopted the Declaration of the Rights of Man and of the Citizen in August 1789.
The monarchy fell in 1792 and the French Republic was proclaimed.
The Reign of Terror (1793-1794) involved emergency government and mass executions.
"""
        cards = compile_study_material(text, "french-revolution.txt")["cards"]
        fronts = {card["front"] for card in cards}
        self.assertGreaterEqual(len(cards), 3)
        self.assertIn("When did the Estates-General meet?", fronts)
        self.assertIn("What did the National Assembly adopt?", fronts)
        self.assertTrue(all(card["source"] == "french-revolution.txt" for card in cards))

    def test_course_schedule_with_standalone_dated_rows_maps_typed_calendar_events(self) -> None:
        content = """PSY 241 Learning and Behavior
Fall 2026 Course Syllabus
Course Schedule
September 9, 2026 - Exam 1: Learning and Conditioning
October 7, 2026 - Quiz 2: Observational Learning
November 18, 2026 - Research paper due
December 9, 2026 - Final examination
"""
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "PSY241_Fall_2026_syllabus.txt"
            path.write_text(content, encoding="utf-8")
            summary = source_summary(path, path.name, "auto")

        self.assertEqual(summary["kind"], "syllabus")
        events = summary["calendarEvents"]
        self.assertEqual({event["date"] for event in events}, {
            "2026-09-09", "2026-10-07", "2026-11-18", "2026-12-09"
        })
        self.assertEqual({event["type"] for event in events}, {"exam", "quiz", "assignment"})
        self.assertTrue(all(event["yearSource"] == "schedule row" for event in events))
        self.assertEqual(summary["draftCards"], [])

    def test_answered_docx_assessment_uses_real_question_and_answer_pairs(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "quiz-answer-key.docx"
            document = Document()
            document.add_heading("Quiz 1 Answer Key", level=1)
            document.add_paragraph("Question 1. What happens to the equilibrium position when a reactant is added?")
            document.add_paragraph("Answer: The system shifts toward products to consume some of the added reactant.")
            document.add_paragraph("Question 2. What does the equilibrium constant describe at a given temperature?")
            document.add_paragraph("Answer: It is the ratio of product concentrations to reactant concentrations, each raised to their stoichiometric coefficients.")
            document.save(path)

            summary = source_summary(path, path.name, "assessment")

        cards = summary["draftCards"]
        self.assertEqual(len(cards), 2)
        self.assertEqual(cards[0]["front"], "What happens to the equilibrium position when a reactant is added?")
        self.assertIn("shifts toward products", cards[0]["back"])
        self.assertEqual(cards[1]["front"], "What does the equilibrium constant describe at a given temperature?")
        self.assertTrue(all(card["field"] == "assessment" for card in cards))
        self.assertTrue(all("Question " in card["source"] for card in cards))
        self.assertFalse(any(card["front"].lower() == "what is answer?" for card in cards))

    def test_pdf_import_does_not_turn_an_in_class_annotation_into_a_study_card(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "lecture-annotations.pdf"
            write_text_pdf(path, ["Boot sequence\nMore detailed (in class): ROM, POST, MBR."])
            summary = source_summary(path, path.name, "material")

        self.assertEqual(summary["draftCards"], [])


if __name__ == "__main__":
    unittest.main()
